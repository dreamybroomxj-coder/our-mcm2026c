"""问题一：严格按用户2026-09-10截图中的MILP模型求解。

运行：python q1.py
依赖：numpy、scipy、openpyxl；不依赖其余问题的代码或商业求解器。
单位：输入功率为 kW，决策变量为每10分钟的电量 kWh，费用为元。
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
from pathlib import Path

import numpy as np
import openpyxl
import scipy
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import coo_matrix

ROOT = Path(__file__).resolve().parent
STEP_HOURS = 1 / 6
CAPACITY = 12000.0
SOC_MIN, SOC_MAX = 0.1 * CAPACITY, 0.9 * CAPACITY
SOC_START = SOC_END = 0.5 * CAPACITY
POWER_MAX = 5000.0
ENERGY_MAX = POWER_MAX * STEP_HOURS
# 主解释：充电效率、放电效率分别为90%，往返效率为81%。
ETA_C = ETA_D = 0.9


def interval_label(slot: int, length: int = 1) -> str:
    """slot 从0起；最后一个时段是23:50-24:00。"""
    start, end = slot * 10, (slot + length) * 10
    return f"{start // 60:02d}:{start % 60:02d}-{end // 60:02d}:{end % 60:02d}"


def right_end_minute(value) -> int:
    if isinstance(value, dt.time):
        return 60 * value.hour + value.minute
    text = str(value).strip()
    if text in ("0:00+1", "00:00+1", "24:00"):
        return 1440
    hour, minute = text.split(":")[:2]
    return 60 * int(hour) + int(minute)


def read_data(path: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """附件1首行为表头，随后144行：时间、电价、负载功率、光伏功率。"""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(book.worksheets[0].values)
    finally:
        book.close()
    if len(rows) != 145 or any(len(row) != 4 for row in rows):
        raise ValueError("附件1应包含1行表头、144行数据、4列。")
    if list(rows[0]) != ["时间", "电价", "小区负载", "光伏发电预测功率"]:
        raise ValueError("附件1列名或列顺序与题目原文件不同，请先核对。")
    minutes = [right_end_minute(row[0]) for row in rows[1:]]
    if minutes != list(range(10, 1441, 10)):
        raise ValueError("时间应为00:10至次日00:00，每10分钟一条。")
    data = np.asarray([row[1:] for row in rows[1:]], dtype=float)
    if not np.isfinite(data).all() or (data < 0).any() or (data[:, 0] <= 0).any():
        raise ValueError("数据必须完整、有限、非负，且电价严格为正。")
    price = data[:, 0]
    # 00:10记录对应00:00—00:10，按区间平均功率换算电量。
    load = data[:, 1] * STEP_HOURS
    pv = data[:, 2] * STEP_HOURS
    return load, pv, price


def solve(load, pv, price):
    """g、c、d、r是微网侧kWh；s是电池内部kWh；u为充放电开关。"""
    n = len(load)
    # 决策向量x=[g, c, d, r, s, u]，各块有n项；s块对应S_1至S_144。
    # 图中r仅代表弃光，不包含已购电弃用。
    rows, cols, values = [], [], []

    def add(row, col, value):
        rows.append(row)
        cols.append(col)
        values.append(value)

    for t in range(n):
        # 图中供需平衡：g + G*Δt - r + d = L*Δt + c。
        for block, coefficient in [(0, 1), (1, -1), (2, 1), (3, -1)]:
            add(t, block * n + t, coefficient)
        # 储电递推：s[t+1] = s[t] + ETA_C*c[t] - d[t]/ETA_D。
        add(n + t, n + t, -ETA_C)
        add(n + t, 2 * n + t, 1 / ETA_D)
        add(n + t, 4 * n + t, 1)
        if t > 0:
            add(n + t, 4 * n + t - 1, -1)
    a_eq = coo_matrix((values, (rows, cols)), shape=(2 * n, 6 * n)).tocsc()
    b_eq = np.r_[load - pv, SOC_START, np.zeros(n - 1)]

    idx = np.arange(n)
    # 目标严格为sum(p*g)，不添加充放电平局成本或终端奖励。
    objective = np.zeros(6 * n)
    objective[:n] = price
    lower = np.zeros(6 * n)
    upper = np.full(6 * n, np.inf)
    upper[n:2 * n] = ENERGY_MAX
    upper[2 * n:3 * n] = ENERGY_MAX
    upper[3 * n:4 * n] = pv
    lower[4 * n:5 * n] = SOC_MIN
    upper[4 * n:5 * n] = SOC_MAX
    lower[5 * n - 1] = upper[5 * n - 1] = SOC_END
    upper[5 * n:6 * n] = 1
    # u=1允许充电，u=0允许放电；允许两者都为0。
    charge_direction = coo_matrix(
        (np.r_[np.ones(n), -np.full(n, ENERGY_MAX)],
         (np.r_[idx, idx], np.r_[n + idx, 5 * n + idx])),
        shape=(n, 6 * n),
    ).tocsc()
    discharge_direction = coo_matrix(
        (np.r_[np.ones(n), np.full(n, ENERGY_MAX)],
         (np.r_[idx, idx], np.r_[2 * n + idx, 5 * n + idx])),
        shape=(n, 6 * n),
    ).tocsc()
    result = milp(
        objective, integrality=np.r_[np.zeros(5 * n), np.ones(n)],
        bounds=Bounds(lower, upper),
        constraints=[LinearConstraint(a_eq, b_eq, b_eq),
                     LinearConstraint(charge_direction, -np.inf, 0),
                     LinearConstraint(discharge_direction, -np.inf, ENERGY_MAX)],
        options={"mip_rel_gap": 1e-9, "time_limit": 60},
    )
    if not result.success:
        raise RuntimeError(f"求解失败：{result.message}")
    # 保留原始求解器数值；不做改变调度的后处理。
    x = result.x.copy()
    solution = {key: x[k * n:(k + 1) * n] for k, key in enumerate(("g", "c", "d", "r", "s", "u"))}
    solution["s"] = np.r_[SOC_START, solution["s"]]
    solution.update(cost=float(price @ solution["g"]), mip_gap=float(result.mip_gap),
                    mip_dual_bound=float(result.mip_dual_bound), solver_message=result.message)
    return solution


def validate(solution, load, pv):
    """从返回轨迹重新核对物理约束；有违反时直接报错。"""
    g, c, d, r, s, u = [solution[k] for k in ("g", "c", "d", "r", "s", "u")]
    checks = {
        "balance_error_kwh": float(np.max(abs(g + pv - r + d - load - c))),
        "state_error_kwh": float(np.max(abs(np.diff(s) - ETA_C * c + d / ETA_D))),
        "negative_flow_violation_kwh": float(max(0, -min(g.min(), c.min(), d.min(), r.min()))),
        "soc_lower_violation_kwh": float(max(0, SOC_MIN - s.min())),
        "soc_upper_violation_kwh": float(max(0, s.max() - SOC_MAX)),
        "charge_limit_violation_kwh": float(max(0, c.max() - ENERGY_MAX)),
        "discharge_limit_violation_kwh": float(max(0, d.max() - ENERGY_MAX)),
        "pv_curtailment_violation_kwh": float(max(0, np.max(r - pv))),
        "charge_binary_violation_kwh": float(max(0, np.max(c - ENERGY_MAX * u))),
        "discharge_binary_violation_kwh": float(max(0, np.max(d - ENERGY_MAX * (1 - u)))),
        "binary_integrality_error": float(np.max(abs(u - np.round(u)))),
        "binary_bound_violation": float(max(0, -u.min(), u.max() - 1)),
        "start_end_error_kwh": float(max(abs(s[0] - SOC_START), abs(s[-1] - SOC_END))),
        "simultaneous_slots": int(np.sum((c > 1e-7) & (d > 1e-7))),
    }
    if max(checks.values()) > 1e-5:
        raise RuntimeError(f"物理约束检查失败：{checks}")
    return checks


def write_outputs(output, template, solution, load, pv, price, summary):
    output.mkdir(parents=True, exist_ok=True)
    g, c, d, r, s, u = [solution[k] for k in ("g", "c", "d", "r", "s", "u")]
    with (output / "q1_detail.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream)
        writer.writerow(["时间段", "电价_元每kWh", "负载_kWh", "光伏_kWh", "购电_kWh",
                         "充电_kWh", "放电_kWh", "弃光_kWh", "段初储电_kWh", "段末储电_kWh", "费用_元", "充放电开关_u"])
        for t in range(len(g)):
            writer.writerow([interval_label(t), price[t], load[t], pv[t], g[t], c[t],
                             d[t], r[t], s[t], s[t + 1], price[t] * g[t], u[t]])
    (output / "q1_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")

    # 在题目模板的副本填写，源模板保持不变；总计写数值，无需公式重算。
    book = openpyxl.load_workbook(template)
    plan_sheet, storage_sheet = book["计划购电量"], book["充放电量"]
    for t in range(len(g)):
        plan_sheet.cell(t + 2, 1, interval_label(t))
        plan_sheet.cell(t + 2, 2, float(g[t]))
    for block in range(6):
        lo, hi = block * 24, (block + 1) * 24
        storage_sheet.cell(block + 2, 1, interval_label(lo, 24))
        storage_sheet.cell(block + 2, 2, float(c[lo:hi].sum()))
        storage_sheet.cell(block + 2, 3, float(d[lo:hi].sum()))
    storage_sheet["D2"], storage_sheet["D3"] = "00:00", "24:00"
    storage_sheet["E2"], storage_sheet["E3"] = float(s[0]), float(s[-1])
    totals = book.create_sheet("结果汇总")
    totals.append(["指标", "数值"])
    for name, value in [
        ("全天购电量（kWh）", float(g.sum())), ("全天购电费用（元）", solution["cost"]),
        ("无储能费用（元）", summary["no_storage_cost_yuan"]),
        ("费用节约比例", summary["saving_fraction"]),
        ("全天充电量（kWh，微网侧）", float(c.sum())),
        ("全天放电量（kWh，微网侧）", float(d.sum())),
        ("全天弃光量（kWh）", float(r.sum())), ("MILP相对最优间隙", solution["mip_gap"]),
    ]:
        totals.append([name, value])
    totals["B3"].number_format = totals["B4"].number_format = "#,##0.00"
    totals["B5"].number_format = "0.00%"
    notes = book.create_sheet("口径说明")
    for row in [
        ["项目", "说明"],
        ["时间对应", "附件00:10对应00:00-00:10；输出副本修正原模板偏移，源文件未修改。"],
        ["供电路径", "外网和光伏可直接供负载，也可充电；容量只约束当时电池储电。"],
        ["效率", "主解释为充电、放电各90%；充放电量为微网侧电量。"],
        ["储能约束", "额定12000kWh，运行范围1200至10800kWh，日初日末均6000kWh。"],
        ["功率约束", "充放电各不超过5000kW，即每10分钟各不超过833.3333kWh。"],
        ["弃光约束", "严格按截图：0≤r≤GΔt，仅允许弃光；没有已购电弃用变量，不售电。"],
        ["求解模型", "纯购电费用目标的MILP；u为二元开关；没有额外平局成本或附加放电负载上限。"],
        ["显示精度", "电量显示4位小数、费用2位小数；底层数值不截断，总计直接保存数值。"],
    ]:
        notes.append(row)
    for sheet in book:
        sheet.freeze_panes = "A2"
        for row in sheet:
            for cell in row:
                cell.font = Font(name="宋体", size=11, bold=cell.row == 1)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="DCE6F1")
                if isinstance(cell.value, (int, float)) and cell.number_format == "General":
                    cell.number_format = "#,##0.0000"
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 23
        if sheet.title in ("结果汇总", "口径说明"):
            sheet.column_dimensions["A"].width = 36
        if sheet.title == "口径说明":
            sheet.column_dimensions["B"].width = 80
    book.save(output / "result1.xlsx")
    book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=ROOT / "input/附件1.xlsx")
    parser.add_argument("--template", type=Path, default=ROOT / "input/result1模板.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT / "results")
    args = parser.parse_args()
    # 防止用户把输出目录指向输入目录而意外覆盖任何输入文件。
    output_files = [args.output / name for name in ("result1.xlsx", "q1_detail.csv", "q1_summary.json")]
    if any(p.resolve() in (args.input.resolve(), args.template.resolve()) for p in output_files):
        raise ValueError("输出文件与输入或模板重合，请指定其他输出目录。")
    load, pv, price = read_data(args.input)
    solution = solve(load, pv, price)
    checks = validate(solution, load, pv)
    baseline = float(price @ np.maximum(load - pv, 0))
    summary = {
        "model": "用户截图中的纯购电费用目标MILP",
        "cost_yuan": solution["cost"], "purchase_kwh": float(solution["g"].sum()),
        "no_storage_cost_yuan": baseline, "saving_fraction": 1 - solution["cost"] / baseline,
        "charge_kwh": float(solution["c"].sum()), "discharge_kwh": float(solution["d"].sum()),
        "initial_soc_kwh": float(solution["s"][0]), "final_soc_kwh": float(solution["s"][-1]),
        "min_soc_kwh": float(solution["s"].min()), "max_soc_kwh": float(solution["s"].max()),
        "pv_curtailment_kwh": float(solution["r"].sum()),
        "mip_gap": solution["mip_gap"], "mip_dual_bound_yuan": solution["mip_dual_bound"],
        "solver_message": solution["solver_message"],
        "eta_charge": ETA_C, "eta_discharge": ETA_D, "checks": checks,
        "numpy_version": np.__version__, "scipy_version": scipy.__version__,
        "input_sha256": hashlib.sha256(args.input.read_bytes()).hexdigest(),
    }
    write_outputs(args.output, args.template, solution, load, pv, price, summary)
    print(f"问题一购电费用：{solution['cost']:,.2f} 元")
    print(f"全天购电量：{summary['purchase_kwh']:,.4f} kWh")
    print(f"无储能费用：{baseline:,.2f} 元；节约：{summary['saving_fraction']:.2%}")
    print("物理约束检查：通过")
    print(f"弃光量：{summary['pv_curtailment_kwh']:.4f} kWh；MILP相对最优间隙：{summary['mip_gap']:.2e}")
    print(f"结果目录：{args.output.resolve()}")


if __name__ == "__main__":
    main()

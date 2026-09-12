"""生成问题三 0/6/12/18 时发布的十分钟区间光伏预测。

候选方案：
A：当前实测边界 + 最新小时预报的分段线性插值；
B：A + 过去 28 个完整日的小时内历史中位形状残差（整点修正为 0）；
C：根据发布前最近 3 小时“旧预报-实测”残差，对最新小时节点作指数衰减修正。

最终方案每天、每个发布时间只使用此前 28 天已经实现的候选误差滚动选优；
历史不足 7 天时回退到 A。正式输出从 2025-02-01 至 2025-12-31。

参考工作簿只读取第一行表头，绝不读取其数据区。

从项目根目录运行：
    python 附件/question3/build_pv_interval_forecasts.py
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import platform
from collections import defaultdict, deque
from dataclasses import dataclass
from datetime import time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_ACTUAL = PROJECT_ROOT / "附件" / "附件2.xlsx"
DEFAULT_FORECAST = PROJECT_ROOT / "附件" / "附件3.xlsx"
DEFAULT_REFERENCE = PROJECT_ROOT / "演示数据用于参考格式.xlsx"
DEFAULT_OUTPUT = Path(__file__).resolve().parent / "问题三_光伏区间预测.xlsx"
DEFAULT_RESULTS_DIR = Path(__file__).resolve().parent / "forecast_results"
OUTPUT_START = pd.Timestamp("2025-02-01")
OUTPUT_END = pd.Timestamp("2025-12-31")
CALCULATION_START = pd.Timestamp("2025-01-08")
RELEASE_HOURS = (0, 6, 12, 18)
METHODS = ("A_线性插值", "B_历史形状", "C_残差同化")
SHAPE_DAYS = 28
SELECTION_DAYS = 28
MIN_SELECTION_DAYS = 7
RECENT_HOURS = 3
BIAS_DECAY_HOURS = 3.0
EXTRA_ENDPOINT_HISTORY_DAYS = 7


@dataclass(frozen=True)
class Inputs:
    actual: dict[pd.Timestamp, float]
    forecasts: dict[pd.Timestamp, np.ndarray]
    reference_headers: list[object]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def parse_header_minute(value: object) -> int:
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    if str(value).strip() in {"0:00+1", "00:00+1", "24:00"}:
        return 1440
    raise ValueError(f"无法识别附件 2 时间表头：{value!r}")


def read_actual(path: Path) -> dict[pd.Timestamp, float]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[1]
    shape = (worksheet.max_row, worksheet.max_column)
    workbook.close()
    if shape != (366, 145):
        raise ValueError(f"附件 2 光伏表应为 366×145，实际为 {shape}")

    frame = pd.read_excel(path, sheet_name=1, header=None)
    minutes = [parse_header_minute(value) for value in frame.iloc[0, 1:]]
    if minutes != list(range(10, 1441, 10)):
        raise ValueError("附件 2 光伏端点不是从 00:10 到次日 00:00 的连续十分钟序列")

    actual: dict[pd.Timestamp, float] = {}
    for row_index in range(1, len(frame)):
        date = pd.Timestamp(frame.iloc[row_index, 0]).normalize()
        values = pd.to_numeric(frame.iloc[row_index, 1:], errors="raise").to_numpy(float)
        if values.shape != (144,) or not np.isfinite(values).all() or np.any(values < 0):
            raise ValueError(f"附件 2 第 {row_index + 1} 行光伏数据异常")
        for minute, value in zip(minutes, values, strict=True):
            actual[date + pd.Timedelta(minutes=minute)] = float(value)
    return actual


def read_forecasts(path: Path) -> dict[pd.Timestamp, np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    shape = (worksheet.max_row, worksheet.max_column)
    workbook.close()
    if shape != (1461, 26):
        raise ValueError(f"附件 3 应为 1461×26，实际为 {shape}")

    frame = pd.read_excel(path, sheet_name=0, header=None)
    dates = frame.iloc[1:, 0].replace("", np.nan).ffill()
    releases = frame.iloc[1:, 1].astype(str).str.strip()
    values = frame.iloc[1:, 2:].apply(pd.to_numeric, errors="raise").to_numpy(float)
    if values.shape != (1460, 24) or not np.isfinite(values).all() or np.any(values < 0):
        raise ValueError("附件 3 预测矩阵异常")

    forecasts: dict[pd.Timestamp, np.ndarray] = {}
    for index, (date_value, release_value) in enumerate(zip(dates, releases, strict=True)):
        date = pd.Timestamp(str(date_value)).normalize()
        release_hour = int(release_value.split(":", maxsplit=1)[0])
        if release_hour not in RELEASE_HOURS:
            raise ValueError(f"未知发布时间：{release_value}")
        issue = date + pd.Timedelta(hours=release_hour)
        forecasts[issue] = values[index].copy()
    if len(forecasts) != 1460:
        raise ValueError(f"附件 3 发布时间数量应为 1460，实际为 {len(forecasts)}")
    return forecasts


def read_reference_headers_only(path: Path) -> list[object]:
    """只读取参考文件第一张工作表第一行，禁止访问第二行及以后。"""
    workbook = load_workbook(path, read_only=True, data_only=False)
    worksheet = workbook.worksheets[0]
    headers = list(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    workbook.close()
    if len(headers) != 146 or headers[0] != "日期\\区间":
        raise ValueError("参考表第一行应为日期列加 145 个区间表头")
    if headers[-1] != "次日00:00—次日00:10":
        raise ValueError("参考表最后一列表头不是次日00:00—次日00:10")
    return headers


def load_inputs(actual_path: Path, forecast_path: Path, reference_path: Path) -> Inputs:
    return Inputs(
        actual=read_actual(actual_path),
        forecasts=read_forecasts(forecast_path),
        reference_headers=read_reference_headers_only(reference_path),
    )


def actual_value(actual: dict[pd.Timestamp, float], timestamp: pd.Timestamp) -> float:
    value = actual.get(timestamp)
    if value is None:
        raise KeyError(f"缺少真实光伏端点：{timestamp}")
    return float(value)


def linear_endpoints(
    issue: pd.Timestamp,
    forecasts: dict[pd.Timestamp, np.ndarray],
    actual: dict[pd.Timestamp, float],
) -> np.ndarray:
    """方案 A：生成 issue 至当天 24:00 的十分钟瞬时端点。"""
    release_hour = issue.hour
    horizon_hours = 24 - release_hour
    hourly_knots = np.concatenate(
        ([actual_value(actual, issue)], forecasts[issue][:horizon_hours])
    )
    endpoint_count = horizon_hours * 6
    endpoints = np.empty(endpoint_count + 1, dtype=float)
    endpoints[0] = hourly_knots[0]
    for index in range(1, endpoint_count + 1):
        right_hour = math.ceil(index / 6)
        offset = index - 6 * (right_hour - 1)
        ratio = offset / 6.0
        endpoints[index] = (
            (1.0 - ratio) * hourly_knots[right_hour - 1]
            + ratio * hourly_knots[right_hour]
        )
    return np.maximum(endpoints, 0.0)


def historical_shape_endpoints(
    issue: pd.Timestamp,
    base_endpoints: np.ndarray,
    actual: dict[pd.Timestamp, float],
) -> np.ndarray:
    """方案 B：用此前完整日的小时内残差修正 A，整点不变。"""
    corrected = base_endpoints.copy()
    history_dates = [
        issue.normalize() - pd.Timedelta(days=lag)
        for lag in range(1, SHAPE_DAYS + 1)
    ]
    for index in range(1, len(corrected)):
        if index % 6 == 0:
            continue
        target = issue + pd.Timedelta(minutes=10 * index)
        hour_start = target.floor("h")
        fraction = target.minute / 60.0
        residuals = []
        for history_date in history_dates:
            historical_start = history_date + pd.Timedelta(hours=hour_start.hour)
            historical_target = historical_start + pd.Timedelta(minutes=target.minute)
            historical_end = historical_start + pd.Timedelta(hours=1)
            if (
                historical_start not in actual
                or historical_target not in actual
                or historical_end not in actual
            ):
                continue
            straight = (
                (1.0 - fraction) * actual[historical_start]
                + fraction * actual[historical_end]
            )
            residuals.append(actual[historical_target] - straight)
        if residuals:
            corrected[index] += float(np.median(residuals))
    return np.maximum(corrected, 0.0)


def previous_forecast_endpoint(
    previous_issue: pd.Timestamp,
    target: pd.Timestamp,
    forecasts: dict[pd.Timestamp, np.ndarray],
    actual: dict[pd.Timestamp, float],
) -> float:
    minutes = int((target - previous_issue).total_seconds() // 60)
    if minutes < 0 or minutes > 24 * 60 or minutes % 10 != 0:
        raise ValueError("旧预报目标时刻超出可插值范围")
    if minutes == 0:
        return actual_value(actual, previous_issue)
    right_hour = math.ceil(minutes / 60)
    offset = minutes - 60 * (right_hour - 1)
    ratio = offset / 60.0
    left = (
        actual_value(actual, previous_issue)
        if right_hour == 1
        else float(forecasts[previous_issue][right_hour - 2])
    )
    right = float(forecasts[previous_issue][right_hour - 1])
    return (1.0 - ratio) * left + ratio * right


def recent_bias(
    issue: pd.Timestamp,
    forecasts: dict[pd.Timestamp, np.ndarray],
    actual: dict[pd.Timestamp, float],
) -> float:
    """根据发布前最近 3 小时旧预报的日间端点残差估计当前偏差。"""
    old_issue = issue - pd.Timedelta(hours=6)
    residuals = []
    for step in range(RECENT_HOURS * 6 + 1):
        target = issue - pd.Timedelta(minutes=10 * (RECENT_HOURS * 6 - step))
        truth = actual_value(actual, target)
        prediction = previous_forecast_endpoint(old_issue, target, forecasts, actual)
        if truth > 1.0 or prediction > 1.0:
            residuals.append(truth - prediction)
    return float(np.median(residuals)) if len(residuals) >= 6 else 0.0


def assimilated_endpoints(
    issue: pd.Timestamp,
    base_endpoints: np.ndarray,
    forecasts: dict[pd.Timestamp, np.ndarray],
    actual: dict[pd.Timestamp, float],
) -> np.ndarray:
    """方案 C：将近期残差以指数衰减形式作用于未来端点。"""
    bias = recent_bias(issue, forecasts, actual)
    corrected = base_endpoints.copy()
    for index in range(1, len(corrected)):
        # 附件 3 已明确给出夜间零功率；不能把白天残差外推成虚假夜间发电。
        if base_endpoints[index] <= 1.0:
            continue
        lead_hours = index / 6.0
        corrected[index] += bias * math.exp(-lead_hours / BIAS_DECAY_HOURS)
    return np.maximum(corrected, 0.0)


def endpoints_to_intervals(endpoints: np.ndarray) -> np.ndarray:
    intervals = (endpoints[:-1] + endpoints[1:]) / 2.0
    if not np.isfinite(intervals).all() or np.any(intervals < 0):
        raise ValueError("端点转换后的区间平均功率异常")
    return intervals


def candidate_intervals(issue: pd.Timestamp, inputs: Inputs) -> dict[str, np.ndarray]:
    base = linear_endpoints(issue, inputs.forecasts, inputs.actual)
    shaped = historical_shape_endpoints(issue, base, inputs.actual)
    assimilated = assimilated_endpoints(issue, base, inputs.forecasts, inputs.actual)
    return {
        "A_线性插值": endpoints_to_intervals(base),
        "B_历史形状": endpoints_to_intervals(shaped),
        "C_残差同化": endpoints_to_intervals(assimilated),
    }


def true_intervals(issue: pd.Timestamp, actual: dict[pd.Timestamp, float]) -> np.ndarray:
    endpoint_count = (24 - issue.hour) * 6
    endpoints = np.array(
        [actual_value(actual, issue + pd.Timedelta(minutes=10 * index)) for index in range(endpoint_count + 1)],
        dtype=float,
    )
    return endpoints_to_intervals(endpoints)


def extra_next_day_interval(
    date: pd.Timestamp,
    left_endpoint: float,
    actual: dict[pd.Timestamp, float],
) -> float:
    """用发布日前 7 个完整日的 00:10 真实端点均值补次日 00:10。"""
    history = [
        actual_value(actual, date - pd.Timedelta(days=lag) + pd.Timedelta(minutes=10))
        for lag in range(1, EXTRA_ENDPOINT_HISTORY_DAYS + 1)
    ]
    next_day_0010 = float(np.mean(history))
    return max(0.0, (left_endpoint + next_day_0010) / 2.0)


def rolling_build(inputs: Inputs) -> tuple[dict[int, list[np.ndarray]], pd.DataFrame, pd.DataFrame]:
    """无泄漏地滚动选择 A/B/C，并返回四个发布时间的正式输出。"""
    error_history: dict[int, dict[str, deque[float]]] = {
        release: {method: deque(maxlen=SELECTION_DAYS) for method in METHODS}
        for release in RELEASE_HOURS
    }
    outputs: dict[int, list[np.ndarray]] = defaultdict(list)
    metric_records: list[dict[str, object]] = []
    choice_records: list[dict[str, object]] = []

    for date in pd.date_range(CALCULATION_START, OUTPUT_END, freq="D"):
        for release_hour in RELEASE_HOURS:
            issue = date + pd.Timedelta(hours=release_hour)
            candidates = candidate_intervals(issue, inputs)
            truth = true_intervals(issue, inputs.actual)
            history = error_history[release_hour]
            history_size = len(history[METHODS[0]])
            history_means_before = {
                method: float(np.mean(history[method])) if history_size else np.nan
                for method in METHODS
            }
            if history_size < MIN_SELECTION_DAYS:
                selected = METHODS[0]
            else:
                selected = min(METHODS, key=lambda name: history_means_before[name])

            daily_errors = {
                method: float(np.mean(np.abs(values - truth)))
                for method, values in candidates.items()
            }
            choice_records.append(
                {
                    "日期": date,
                    "发布时间": release_hour,
                    "选择方法": selected,
                    "选择依据历史天数": history_size,
                    "选择前历史MAE_A_kW": history_means_before[METHODS[0]],
                    "选择前历史MAE_B_kW": history_means_before[METHODS[1]],
                    "选择前历史MAE_C_kW": history_means_before[METHODS[2]],
                }
            )
            for method, error in daily_errors.items():
                history[method].append(error)
                metric_records.append(
                    {
                        "日期": date,
                        "发布时间": release_hour,
                        "方法": method,
                        "当日有效区间MAE_kW": error,
                    }
                )

            if OUTPUT_START <= date <= OUTPUT_END:
                chosen = candidates[selected]
                final_hour_endpoint = float(inputs.forecasts[issue][24 - release_hour - 1])
                extra = extra_next_day_interval(date, final_hour_endpoint, inputs.actual)
                output_values = np.concatenate((chosen, [extra]))
                expected = 145 - 6 * release_hour
                if len(output_values) != expected:
                    raise AssertionError(
                        f"{date.date()} {release_hour}:00 应输出 {expected} 个区间，实际 {len(output_values)}"
                    )
                outputs[release_hour].append(output_values)

    return outputs, pd.DataFrame(metric_records), pd.DataFrame(choice_records)


def write_forecast_workbook(
    output_path: Path,
    headers: list[object],
    outputs: dict[int, list[np.ndarray]],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    dates = list(pd.date_range(OUTPUT_START, OUTPUT_END, freq="D"))
    fill = PatternFill("solid", fgColor="D9EAF7")

    for release_hour in RELEASE_HOURS:
        worksheet = workbook.create_sheet(f"{release_hour:02d}时预报")
        start_interval = release_hour * 6
        sheet_headers = [headers[0], *headers[1 + start_interval :]]
        worksheet.append(sheet_headers)
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        if len(outputs[release_hour]) != len(dates):
            raise AssertionError(f"{release_hour}:00 输出日期数异常")
        for date, values in zip(dates, outputs[release_hour], strict=True):
            worksheet.append([date.to_pydatetime(), *[float(value) for value in values]])
        worksheet.freeze_panes = "B2"
        worksheet.column_dimensions["A"].width = 13
        for cell in worksheet["A"][1:]:
            cell.number_format = "yyyy-mm-dd"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_supporting_results(
    results_dir: Path,
    metrics: pd.DataFrame,
    choices: pd.DataFrame,
    inputs: Inputs,
    actual_path: Path,
    forecast_path: Path,
    reference_path: Path,
    output_path: Path,
) -> None:
    results_dir.mkdir(parents=True, exist_ok=True)
    formal_metrics = metrics[metrics["日期"] >= OUTPUT_START].copy()
    summary = (
        formal_metrics.groupby(["发布时间", "方法"], sort=True)["当日有效区间MAE_kW"]
        .agg(["count", "mean", "median"])
        .reset_index()
        .rename(columns={"count": "天数", "mean": "日MAE均值_kW", "median": "日MAE中位数_kW"})
    )
    formal_choices = choices[choices["日期"] >= OUTPUT_START].copy()
    choice_summary = (
        formal_choices.groupby(["发布时间", "选择方法"], sort=True)
        .size()
        .reset_index(name="选择天数")
    )
    summary.to_csv(results_dir / "ABC方案误差汇总.csv", index=False, encoding="utf-8-sig")
    formal_choices.to_csv(results_dir / "每日滚动方法选择.csv", index=False, encoding="utf-8-sig")
    choice_summary.to_csv(results_dir / "方法选择次数.csv", index=False, encoding="utf-8-sig")

    manifest = {
        "purpose": "问题三光伏小时预报还原为十分钟区间预测",
        "output_period": [str(OUTPUT_START.date()), str(OUTPUT_END.date())],
        "methods": list(METHODS),
        "selection": f"同发布时间此前{SELECTION_DAYS}天滚动日MAE最小；不足{MIN_SELECTION_DAYS}天用A",
        "historical_shape_days": SHAPE_DAYS,
        "recent_residual_hours": RECENT_HOURS,
        "bias_decay_hours": BIAS_DECAY_HOURS,
        "extra_0010_rule": "发布日前7个完整日的00:10真实端点平均，再与24:00预测端点取平均",
        "reference_access": "只读取演示数据用于参考格式.xlsx第一张表第一行表头",
        "inputs": {
            str(actual_path.relative_to(PROJECT_ROOT)): sha256(actual_path),
            str(forecast_path.relative_to(PROJECT_ROOT)): sha256(forecast_path),
        },
        "reference_header_source": str(reference_path.relative_to(PROJECT_ROOT)),
        "program": {
            str(Path(__file__).resolve().relative_to(PROJECT_ROOT)): sha256(Path(__file__).resolve())
        },
        "output": {
            str(output_path.relative_to(PROJECT_ROOT)): sha256(output_path),
            "sheet_shapes": {
                f"{release:02d}时预报": [335, 146 - 6 * release]
                for release in RELEASE_HOURS
            },
        },
        "runtime": {
            "python": platform.python_version(),
            "numpy": np.__version__,
            "pandas": pd.__version__,
        },
        "reproduce_from_project_root": "python 附件/question3/build_pv_interval_forecasts.py",
    }
    (results_dir / "复现清单.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def validate_output(path: Path, headers: list[object]) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    expected_sheets = [f"{release:02d}时预报" for release in RELEASE_HOURS]
    if workbook.sheetnames != expected_sheets:
        raise AssertionError(f"工作表名称异常：{workbook.sheetnames}")
    for release_hour, sheet_name in zip(RELEASE_HOURS, expected_sheets, strict=True):
        worksheet = workbook[sheet_name]
        expected_columns = 146 - 6 * release_hour
        if (worksheet.max_row, worksheet.max_column) != (335, expected_columns):
            raise AssertionError(
                f"{sheet_name} 应为 335×{expected_columns}，实际为 "
                f"{worksheet.max_row}×{worksheet.max_column}"
            )
        actual_headers = [cell.value for cell in worksheet[1]]
        expected_headers = [headers[0], *headers[1 + 6 * release_hour :]]
        if actual_headers != expected_headers:
            raise AssertionError(f"{sheet_name} 表头未按参考文件截取")
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))
        dates = [pd.Timestamp(row[0]).normalize() for row in rows]
        if dates[0] != OUTPUT_START or dates[-1] != OUTPUT_END or len(set(dates)) != 334:
            raise AssertionError(f"{sheet_name} 日期范围或唯一性异常")
        values = np.asarray([row[1:] for row in rows], dtype=float)
        if not np.isfinite(values).all() or np.any(values < 0):
            raise AssertionError(f"{sheet_name} 存在缺失、无穷或负功率")
    workbook.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--actual", type=Path, default=DEFAULT_ACTUAL)
    parser.add_argument("--forecast", type=Path, default=DEFAULT_FORECAST)
    parser.add_argument("--reference", type=Path, default=DEFAULT_REFERENCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--results-dir", type=Path, default=DEFAULT_RESULTS_DIR)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    actual_path = args.actual.resolve()
    forecast_path = args.forecast.resolve()
    reference_path = args.reference.resolve()
    output_path = args.output.resolve()
    results_dir = args.results_dir.resolve()
    inputs = load_inputs(actual_path, forecast_path, reference_path)
    outputs, metrics, choices = rolling_build(inputs)
    write_forecast_workbook(output_path, inputs.reference_headers, outputs)
    validate_output(output_path, inputs.reference_headers)
    write_supporting_results(
        results_dir,
        metrics,
        choices,
        inputs,
        actual_path,
        forecast_path,
        reference_path,
        output_path,
    )
    print(f"已生成：{output_path}")
    print("工作表：00时预报、06时预报、12时预报、18时预报")
    print("数据行：每张表 334 天；仅含日期和光伏区间预测")


if __name__ == "__main__":
    main()

"""比较“上一发布时刻延伸预报”和“当前最新预报”的准确性。

比较遵循严格的配对原则：
1. 两版预报必须针对同一个目标时刻；
2. 只评价当前发布时刻至当天 24:00、且旧预报仍能覆盖的区间；
3. 小时端点评价使用附件 2 的同一整点真实功率；
4. 10 分钟区间评价时，两种方案都用当前发布时刻的真实功率作共同边界，
   再分别对旧、新小时节点作线性插值和梯形平均。

从项目根目录运行：
    python 附件/question3/compare_forecast_versions.py
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import platform
from dataclasses import dataclass
from datetime import time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_ACTUAL = PROJECT_ROOT / "附件" / "附件2.xlsx"
DEFAULT_FORECAST = PROJECT_ROOT / "附件" / "附件3.xlsx"
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parent / "results"
START_DATE = pd.Timestamp("2025-02-01")
END_DATE = pd.Timestamp("2025-12-31")
RELEASE_HOURS = (0, 6, 12, 18)


@dataclass(frozen=True)
class InputData:
    actual: dict[pd.Timestamp, float]
    forecasts: dict[pd.Timestamp, np.ndarray]


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def validate_workbook_shape(path: Path, sheet_index: int, rows: int, columns: int) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[sheet_index]
    actual_shape = (worksheet.max_row, worksheet.max_column)
    workbook.close()
    if actual_shape != (rows, columns):
        raise ValueError(
            f"{path.name} 第 {sheet_index + 1} 张表尺寸应为 {(rows, columns)}，"
            f"实际为 {actual_shape}"
        )


def parse_header_minute(value: object) -> int:
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    text = str(value).strip()
    if text in {"0:00+1", "00:00+1", "24:00"}:
        return 1440
    raise ValueError(f"无法识别附件 2 时间表头：{value!r}")


def read_actual_pv(path: Path) -> dict[pd.Timestamp, float]:
    """读取附件 2 第二张表，返回所有 10 分钟真实端点。"""
    validate_workbook_shape(path, sheet_index=1, rows=366, columns=145)
    frame = pd.read_excel(path, sheet_name=1, header=None)
    minutes = [parse_header_minute(value) for value in frame.iloc[0, 1:]]
    if minutes != list(range(10, 1441, 10)):
        raise ValueError("附件 2 光伏表头不是 00:10 至次日 00:00 的连续 10 分钟端点")

    actual: dict[pd.Timestamp, float] = {}
    for row_index in range(1, len(frame)):
        date = pd.Timestamp(frame.iloc[row_index, 0]).normalize()
        values = pd.to_numeric(frame.iloc[row_index, 1:], errors="raise").to_numpy(float)
        if values.shape != (144,) or not np.isfinite(values).all() or np.any(values < 0):
            raise ValueError(f"附件 2 第 {row_index + 1} 行光伏数据异常")
        for minute, value in zip(minutes, values, strict=True):
            actual[date + pd.Timedelta(minutes=minute)] = float(value)

    expected_first = pd.Timestamp("2025-01-01 00:10")
    expected_last = pd.Timestamp("2026-01-01 00:00")
    if expected_first not in actual or expected_last not in actual:
        raise ValueError("附件 2 实际数据首末时间不符合预期")
    return actual


def read_forecasts(path: Path) -> dict[pd.Timestamp, np.ndarray]:
    """读取附件 3，键为完整发布时间，值为未来 1—24 小时端点预测。"""
    validate_workbook_shape(path, sheet_index=0, rows=1461, columns=26)
    frame = pd.read_excel(path, sheet_name=0, header=None)
    dates = frame.iloc[1:, 0].replace("", np.nan).ffill()
    release_text = frame.iloc[1:, 1].astype(str).str.strip()
    values = frame.iloc[1:, 2:].apply(pd.to_numeric, errors="raise").to_numpy(float)
    if values.shape != (1460, 24) or not np.isfinite(values).all() or np.any(values < 0):
        raise ValueError("附件 3 预测矩阵尺寸或数值异常")

    forecasts: dict[pd.Timestamp, np.ndarray] = {}
    for index, (date_value, release_value) in enumerate(zip(dates, release_text, strict=True)):
        date = pd.Timestamp(str(date_value)).normalize()
        hour = int(release_value.split(":", maxsplit=1)[0])
        if hour not in RELEASE_HOURS:
            raise ValueError(f"附件 3 出现非预期发布时间：{release_value}")
        issue = date + pd.Timedelta(hours=hour)
        if issue in forecasts:
            raise ValueError(f"附件 3 发布时间重复：{issue}")
        forecasts[issue] = values[index].copy()

    if len(forecasts) != 365 * 4:
        raise ValueError(f"附件 3 应有 1460 个发布时间，实际为 {len(forecasts)}")
    return forecasts


def load_inputs(actual_path: Path, forecast_path: Path) -> InputData:
    return InputData(
        actual=read_actual_pv(actual_path),
        forecasts=read_forecasts(forecast_path),
    )


def overlap_hours(release_hour: int) -> int:
    """当前日剩余范围与上一版 24 小时预报覆盖范围的交集小时数。"""
    return min(24 - release_hour, 18)


def previous_issue(issue: pd.Timestamp) -> pd.Timestamp:
    return issue - pd.Timedelta(hours=6)


def collect_hourly_pairs(data: InputData) -> pd.DataFrame:
    records: list[dict[str, object]] = []
    for date in pd.date_range(START_DATE, END_DATE, freq="D"):
        for release_hour in RELEASE_HOURS:
            issue = date + pd.Timedelta(hours=release_hour)
            old_issue = previous_issue(issue)
            if issue not in data.forecasts or old_issue not in data.forecasts:
                raise KeyError(f"缺少发布时间 {issue} 或 {old_issue} 的预测")
            new_values = data.forecasts[issue]
            old_values = data.forecasts[old_issue]
            for horizon in range(1, overlap_hours(release_hour) + 1):
                target = issue + pd.Timedelta(hours=horizon)
                old_lead = horizon + 6
                truth = data.actual.get(target)
                if truth is None:
                    raise KeyError(f"缺少目标时刻 {target} 的实际光伏功率")
                old_prediction = float(old_values[old_lead - 1])
                new_prediction = float(new_values[horizon - 1])
                records.append(
                    {
                        "日期": date,
                        "当前发布时间": issue,
                        "上一发布时间": old_issue,
                        "发布小时": release_hour,
                        "当前预测提前小时": horizon,
                        "旧预测提前小时": old_lead,
                        "目标时刻": target,
                        "真实功率_kW": truth,
                        "旧版延伸预测_kW": old_prediction,
                        "新版最新预测_kW": new_prediction,
                        "旧版绝对误差_kW": abs(old_prediction - truth),
                        "新版绝对误差_kW": abs(new_prediction - truth),
                    }
                )
    result = pd.DataFrame.from_records(records)
    result["绝对误差改善_kW"] = result["旧版绝对误差_kW"] - result["新版绝对误差_kW"]
    return result


def interpolate_endpoint(knots: np.ndarray, ten_minute_index: int) -> float:
    """从逐小时节点线性插值得到发布后的第 ten_minute_index 个端点。"""
    hour_right = math.ceil(ten_minute_index / 6)
    offset = ten_minute_index - 6 * (hour_right - 1)
    ratio = offset / 6.0
    return float((1.0 - ratio) * knots[hour_right - 1] + ratio * knots[hour_right])


def collect_interval_pairs(data: InputData) -> pd.DataFrame:
    """比较共同实测边界下，两版预测重构得到的 10 分钟区间平均功率。"""
    records: list[dict[str, object]] = []
    for date in pd.date_range(START_DATE, END_DATE, freq="D"):
        for release_hour in RELEASE_HOURS:
            issue = date + pd.Timedelta(hours=release_hour)
            old_issue = previous_issue(issue)
            hours = overlap_hours(release_hour)
            boundary = data.actual.get(issue)
            if boundary is None:
                raise KeyError(f"缺少发布边界 {issue} 的实际光伏功率")

            new_hourly = data.forecasts[issue][:hours]
            old_hourly = data.forecasts[old_issue][6 : 6 + hours]
            new_knots = np.concatenate(([boundary], new_hourly))
            old_knots = np.concatenate(([boundary], old_hourly))

            previous_new_endpoint = float(boundary)
            previous_old_endpoint = float(boundary)
            previous_actual_endpoint = float(boundary)
            for ten_minute_index in range(1, hours * 6 + 1):
                target = issue + pd.Timedelta(minutes=10 * ten_minute_index)
                actual_endpoint = data.actual.get(target)
                if actual_endpoint is None:
                    raise KeyError(f"缺少目标时刻 {target} 的实际光伏功率")
                new_endpoint = interpolate_endpoint(new_knots, ten_minute_index)
                old_endpoint = interpolate_endpoint(old_knots, ten_minute_index)
                actual_interval = (previous_actual_endpoint + actual_endpoint) / 2.0
                new_interval = (previous_new_endpoint + new_endpoint) / 2.0
                old_interval = (previous_old_endpoint + old_endpoint) / 2.0
                records.append(
                    {
                        "日期": date,
                        "当前发布时间": issue,
                        "上一发布时间": old_issue,
                        "发布小时": release_hour,
                        "发布后区间序号": ten_minute_index,
                        "当前预测提前小时": math.ceil(ten_minute_index / 6),
                        "区间开始": target - pd.Timedelta(minutes=10),
                        "区间结束": target,
                        "真实区间平均功率_kW": actual_interval,
                        "旧版延伸区间预测_kW": old_interval,
                        "新版最新区间预测_kW": new_interval,
                        "旧版绝对误差_kW": abs(old_interval - actual_interval),
                        "新版绝对误差_kW": abs(new_interval - actual_interval),
                    }
                )
                previous_new_endpoint = new_endpoint
                previous_old_endpoint = old_endpoint
                previous_actual_endpoint = float(actual_endpoint)

    result = pd.DataFrame.from_records(records)
    result["绝对误差改善_kW"] = result["旧版绝对误差_kW"] - result["新版绝对误差_kW"]
    return result


def metrics(group: pd.DataFrame, truth_column: str, old_column: str, new_column: str) -> pd.Series:
    truth = group[truth_column].to_numpy(float)
    old = group[old_column].to_numpy(float)
    new = group[new_column].to_numpy(float)
    old_error = old - truth
    new_error = new - truth
    old_mae = float(np.mean(np.abs(old_error)))
    new_mae = float(np.mean(np.abs(new_error)))
    denominator = float(np.sum(np.abs(truth)))
    old_error_sum = float(np.sum(np.abs(old_error)))
    new_error_sum = float(np.sum(np.abs(new_error)))
    if old_mae == 0.0:
        if new_mae != 0.0:
            raise ValueError("旧版 MAE 为 0 而新版非 0，MAE 相对改善率无有限定义")
        mae_improvement = 0.0
    else:
        mae_improvement = 100.0 * (old_mae - new_mae) / old_mae
    if denominator == 0.0:
        if old_error_sum != 0.0 or new_error_sum != 0.0:
            raise ValueError("真实功率全为 0 但预测误差非 0，WAPE 无有限定义")
        old_wape = 0.0
        new_wape = 0.0
    else:
        old_wape = 100.0 * old_error_sum / denominator
        new_wape = 100.0 * new_error_sum / denominator
    return pd.Series(
        {
            "样本数": len(group),
            "旧版MAE_kW": old_mae,
            "新版MAE_kW": new_mae,
            "MAE改善率_percent": mae_improvement,
            "旧版RMSE_kW": float(np.sqrt(np.mean(old_error**2))),
            "新版RMSE_kW": float(np.sqrt(np.mean(new_error**2))),
            "旧版WAPE_percent": old_wape,
            "新版WAPE_percent": new_wape,
            "旧版Bias_kW": float(np.mean(old_error)),
            "新版Bias_kW": float(np.mean(new_error)),
            "新版绝对误差更小比例_percent": 100.0 * float(np.mean(np.abs(new_error) < np.abs(old_error))),
            "两版绝对误差相同比例_percent": 100.0 * float(np.mean(np.isclose(np.abs(new_error), np.abs(old_error)))),
        }
    )


def summarize(
    frame: pd.DataFrame,
    truth_column: str,
    old_column: str,
    new_column: str,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    overall = metrics(frame, truth_column, old_column, new_column).to_frame().T
    by_release = (
        frame.groupby("发布小时", sort=True)
        .apply(metrics, truth_column, old_column, new_column, include_groups=False)
        .reset_index()
    )
    by_horizon = (
        frame.groupby(["发布小时", "当前预测提前小时"], sort=True)
        .apply(metrics, truth_column, old_column, new_column, include_groups=False)
        .reset_index()
    )
    return overall, by_release, by_horizon


def monthly_summary(
    frame: pd.DataFrame,
    truth_column: str,
    old_column: str,
    new_column: str,
) -> pd.DataFrame:
    local = frame.copy()
    local["月份"] = pd.to_datetime(local["日期"]).dt.month
    return (
        local.groupby(["月份", "发布小时"], sort=True)
        .apply(metrics, truth_column, old_column, new_column, include_groups=False)
        .reset_index()
    )


def conclusion_table(interval_by_release: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, row in interval_by_release.iterrows():
        improvement = float(row["MAE改善率_percent"])
        absolute_difference = float(row["旧版MAE_kW"] - row["新版MAE_kW"])
        if improvement >= 0.5:
            comparison = "新版明显更优"
            recommendation = "只用当前最新预报"
        elif abs(improvement) < 0.5 or abs(absolute_difference) < 1.0:
            comparison = "两版精度近似等价"
            recommendation = "为保持规则统一和流程简洁，只用当前最新预报"
        else:
            comparison = "旧版明显更优"
            recommendation = "暂不舍弃旧版，进一步研究组合或分时选择"
        rows.append(
            {
                "发布时间": f"{int(row['发布小时']):02d}:00",
                "旧版区间MAE_kW": row["旧版MAE_kW"],
                "新版区间MAE_kW": row["新版MAE_kW"],
                "新版MAE改善率_percent": improvement,
                "精度判断": comparison,
                "建议": recommendation,
            }
        )
    return pd.DataFrame(rows)


def write_conclusion_markdown(path: Path, decisions: pd.DataFrame) -> None:
    lines = [
        "# 延伸预报与最新预报对比结论",
        "",
        "比较采用同一目标时刻配对，并且只评价当前发布时刻至当天 24:00、旧版仍可覆盖的共同区间。",
        "10 分钟区间重构时，两版均以当前发布时刻的真实光伏功率为共同边界。",
        "",
        "| 发布时间 | 旧版区间 MAE/kW | 新版区间 MAE/kW | 新版改善率 | 判断 |",
        "|---|---:|---:|---:|---|",
    ]
    for _, row in decisions.iterrows():
        lines.append(
            f"| {row['发布时间']} | {row['旧版区间MAE_kW']:.4f} | "
            f"{row['新版区间MAE_kW']:.4f} | {row['新版MAE改善率_percent']:.2f}% | "
            f"{row['精度判断']} |"
        )
    lines.extend(
        [
            "",
            "## 建议",
            "",
            "当天 0:00 制定计划时，不将前一天 18:00 延伸预报混入中央预测，直接采用当天 0:00 最新预报。",
            "同理，6:00、12:00 和 18:00 调整时均采用当前最新发布版本覆盖尚未执行区间。",
            "旧版预报只保留用于预测修订分析、缺失备用和不确定性场景，不参与主方案插值。",
            "",
            "18:00 后两版区间 MAE 的差异远小于 1 kW，主要因为该时段基本进入夜间，"
            "不能把这种数值级微差解释为旧版具有实际优势。",
            "纯夜间分组若真实值、旧预测和新预测均为 0，则其 WAPE 与 MAE 改善率按约定记为 0，"
            "表示两版在该全零分组中均无误差。",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_outputs(
    output_dir: Path,
    actual_path: Path,
    forecast_path: Path,
    hourly: pd.DataFrame,
    intervals: pd.DataFrame,
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    hourly_overall, hourly_by_release, hourly_by_horizon = summarize(
        hourly,
        "真实功率_kW",
        "旧版延伸预测_kW",
        "新版最新预测_kW",
    )
    interval_overall, interval_by_release, interval_by_horizon = summarize(
        intervals,
        "真实区间平均功率_kW",
        "旧版延伸区间预测_kW",
        "新版最新区间预测_kW",
    )
    hourly_monthly = monthly_summary(
        hourly,
        "真实功率_kW",
        "旧版延伸预测_kW",
        "新版最新预测_kW",
    )
    interval_monthly = monthly_summary(
        intervals,
        "真实区间平均功率_kW",
        "旧版延伸区间预测_kW",
        "新版最新区间预测_kW",
    )
    decisions = conclusion_table(interval_by_release)

    hourly.to_csv(output_dir / "逐整点配对明细.csv", index=False, encoding="utf-8-sig")
    intervals.to_csv(output_dir / "逐10分钟区间配对明细.csv", index=False, encoding="utf-8-sig")
    interval_by_release.to_csv(output_dir / "核心对比_按发布时间.csv", index=False, encoding="utf-8-sig")
    write_conclusion_markdown(output_dir / "结论说明.md", decisions)

    workbook_path = output_dir / "延伸预报与最新预报准确率对比.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        decisions.to_excel(writer, sheet_name="结论", index=False)
        hourly_overall.to_excel(writer, sheet_name="整点_总体", index=False)
        hourly_by_release.to_excel(writer, sheet_name="整点_按发布时间", index=False)
        hourly_by_horizon.to_excel(writer, sheet_name="整点_按时距", index=False)
        hourly_monthly.to_excel(writer, sheet_name="整点_逐月", index=False)
        interval_overall.to_excel(writer, sheet_name="区间_总体", index=False)
        interval_by_release.to_excel(writer, sheet_name="区间_按发布时间", index=False)
        interval_by_horizon.to_excel(writer, sheet_name="区间_按时距", index=False)
        interval_monthly.to_excel(writer, sheet_name="区间_逐月", index=False)

    manifest = {
        "purpose": "同目标时刻配对比较上一发布时刻延伸预报与当前最新预报",
        "evaluation_period": [str(START_DATE.date()), str(END_DATE.date())],
        "release_hours": list(RELEASE_HOURS),
        "comparison_scope": "当前发布时刻至当天24:00，且旧版预报仍可覆盖的共同范围",
        "interval_reconstruction": "当前实测边界锚定、小时节点线性插值、相邻端点梯形平均",
        "inputs": {
            str(actual_path.relative_to(PROJECT_ROOT)): file_sha256(actual_path),
            str(forecast_path.relative_to(PROJECT_ROOT)): file_sha256(forecast_path),
        },
        "program": {
            str(Path(__file__).resolve().relative_to(PROJECT_ROOT)): file_sha256(Path(__file__).resolve()),
        },
        "outputs": {
            "hourly_pairs": len(hourly),
            "ten_minute_interval_pairs": len(intervals),
            "workbook": workbook_path.name,
        },
        "runtime": {
            "python": platform.python_version(),
            "numpy": np.__version__,
            "pandas": pd.__version__,
        },
        "reproduce_from_project_root": "python 附件/question3/compare_forecast_versions.py",
    }
    (output_dir / "复现清单.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--actual", type=Path, default=DEFAULT_ACTUAL)
    parser.add_argument("--forecast", type=Path, default=DEFAULT_FORECAST)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    actual_path = args.actual.resolve()
    forecast_path = args.forecast.resolve()
    output_dir = args.output_dir.resolve()
    data = load_inputs(actual_path, forecast_path)
    hourly = collect_hourly_pairs(data)
    intervals = collect_interval_pairs(data)
    write_outputs(output_dir, actual_path, forecast_path, hourly, intervals)
    print(f"逐整点配对样本数：{len(hourly)}")
    print(f"逐10分钟区间配对样本数：{len(intervals)}")
    print(f"结果目录：{output_dir}")


if __name__ == "__main__":
    main()

"""把现有端点预测转换为144个区间平均预测，并重新计算统一指标。"""

from __future__ import annotations

import argparse
import math
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import r2_score

from forecast_models import LOAD_SHEET, PV_SHEET, read_power_sheet
from interval_preprocessing import interval_average_actual, period_mapping


ORIGINAL_MODELS = ("Baseline", "XGBoost", "LightGBM", "SARIMA")
IMPROVED_PV_MODELS = ("日总量+归一化曲线", "晴空包络短期外推", "自适应加权Baseline")
FIRST_FORECAST_INDEX = 31


def endpoint_to_interval(predicted_endpoints: np.ndarray, observed_00: np.ndarray) -> np.ndarray:
    """真实00:00与预测00:10—次日00:00组成145端点，再取相邻均值。"""
    if predicted_endpoints.ndim != 2 or predicted_endpoints.shape[1] != 144:
        raise ValueError("端点预测必须为 n×144 矩阵")
    if len(observed_00) != len(predicted_endpoints):
        raise ValueError("00:00真实边界长度与预测天数不一致")
    endpoints = np.column_stack([observed_00, predicted_endpoints])
    return (endpoints[:, :-1] + endpoints[:, 1:]) / 2


def basic_metrics(actual: np.ndarray, predicted: np.ndarray, target: str) -> dict[str, float]:
    error = predicted - actual
    result = {
        "MAE": float(np.abs(error).mean()),
        "RMSE": float(np.sqrt(np.square(error).mean())),
        "WAPE_percent": float(np.abs(error).sum() / np.abs(actual).sum() * 100),
        "Bias": float(error.mean()),
        "R2": float(r2_score(actual.ravel(), predicted.ravel())),
    }
    actual_energy = actual.sum(axis=1) / 6
    predicted_energy = predicted.sum(axis=1) / 6
    result["Daily_energy_MAPE_percent"] = float(
        np.mean(np.abs(predicted_energy - actual_energy) / np.maximum(actual_energy, 1e-8)) * 100
    )
    result["Daily_energy_MAE"] = float(np.abs(predicted_energy - actual_energy).mean())
    result["Daily_peak_MAE"] = float(np.abs(predicted.max(axis=1) - actual.max(axis=1)).mean())
    if target == "光伏":
        daylight = actual > 1.0
        result["Daylight_MAE"] = float(np.abs(error[daylight]).mean())
        result["Daylight_WAPE_percent"] = float(
            np.abs(error[daylight]).sum() / actual[daylight].sum() * 100
        )
        result["Overforecast_rate_percent"] = float(
            (predicted[daylight] > actual[daylight]).mean() * 100
        )
    return result


def wide_frame(dates, labels, values):
    frame = pd.DataFrame(values, columns=labels)
    frame.insert(0, "日期\\区间", dates)
    return frame


def save_baseline_pipeline(path, dates, labels, load_prediction, pv_prediction):
    long = pd.DataFrame(
        {
            "日期": np.repeat(dates, 144),
            "时段序号": np.tile(np.arange(1, 145), len(dates)),
            "时间区间": np.tile(labels, len(dates)),
            "负载区间平均预测功率": load_prediction.ravel(),
            "光伏区间平均预测功率": pv_prediction.ravel(),
        }
    )
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        wide_frame(dates, labels, load_prediction).to_excel(writer, sheet_name="负载区间预测", index=False)
        wide_frame(dates, labels, pv_prediction).to_excel(writer, sheet_name="光伏区间预测", index=False)
        long.to_excel(writer, sheet_name="预测长表", index=False)


def save_all_predictions(path, dates, labels, predictions, metrics):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        metrics.to_excel(writer, sheet_name="区间预测指标", index=False)
        for (model, target), values in predictions.items():
            wide_frame(dates, labels, values).to_excel(
                writer, sheet_name=f"{model}_{target}"[:31], index=False
            )


def validate_excel(path: Path, expected_sheets: int) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if len(workbook.sheetnames) != expected_sheets:
        raise AssertionError(f"工作表数量异常：{len(workbook.sheetnames)}")
    for sheet_name in workbook.sheetnames[1:]:
        worksheet = workbook[sheet_name]
        if worksheet.max_row != 335 or worksheet.max_column != 145:
            raise AssertionError(f"{sheet_name}维度异常")
    workbook.close()


def main() -> None:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--attachment2", type=Path, default=script_dir.parent / "附件2.xlsx")
    parser.add_argument("--original-results", type=Path, default=script_dir / "results")
    parser.add_argument("--improved-results", type=Path, default=script_dir / "improved_pv_results")
    parser.add_argument("--output-dir", type=Path, default=script_dir / "interval_results")
    args = parser.parse_args()

    load_dates, _, load_actual_endpoints = read_power_sheet(args.attachment2, LOAD_SHEET)
    pv_dates, _, pv_actual_endpoints = read_power_sheet(args.attachment2, PV_SHEET)
    if load_dates != pv_dates:
        raise ValueError("附件2负载和光伏日期不一致")
    dates = load_dates[FIRST_FORECAST_INDEX:]
    labels = period_mapping()["时间区间"].tolist()

    load_00, load_actual_all = interval_average_actual(load_actual_endpoints)
    pv_00, pv_actual_all = interval_average_actual(pv_actual_endpoints)
    actuals = {
        "负载": load_actual_all[FIRST_FORECAST_INDEX:],
        "光伏": pv_actual_all[FIRST_FORECAST_INDEX:],
    }
    boundaries = {
        "负载": load_00[FIRST_FORECAST_INDEX:],
        "光伏": pv_00[FIRST_FORECAST_INDEX:],
    }

    predictions: dict[tuple[str, str], np.ndarray] = {}
    for model in ORIGINAL_MODELS:
        for target in ("负载", "光伏"):
            endpoint_path = args.original_results / f"{model}_{target}_predictions.npy"
            endpoint_prediction = np.load(endpoint_path)
            predictions[(model, target)] = endpoint_to_interval(
                endpoint_prediction, boundaries[target]
            )
    for model in IMPROVED_PV_MODELS:
        endpoint_path = args.improved_results / f"{model}_predictions.npy"
        predictions[(model, "光伏")] = endpoint_to_interval(
            np.load(endpoint_path), boundaries["光伏"]
        )

    rows = []
    monthly_rows = []
    for (model, target), predicted in predictions.items():
        actual = actuals[target]
        if predicted.shape != (334, 144) or np.any(predicted < 0) or not np.isfinite(predicted).all():
            raise AssertionError(f"{model}-{target}区间预测异常")
        rows.append({"模型": model, "对象": target, **basic_metrics(actual, predicted, target)})
        for month in range(2, 13):
            mask = np.asarray([date.month == month for date in dates])
            monthly_rows.append(
                {"模型": model, "对象": target, "月份": month, **basic_metrics(actual[mask], predicted[mask], target)}
            )

    metrics = pd.DataFrame(rows).sort_values(["对象", "WAPE_percent"]).reset_index(drop=True)
    monthly_metrics = pd.DataFrame(monthly_rows)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    metrics.to_csv(args.output_dir / "区间预测综合指标.csv", index=False, encoding="utf-8-sig")
    monthly_metrics.to_csv(args.output_dir / "区间预测逐月指标.csv", index=False, encoding="utf-8-sig")
    for (model, target), values in predictions.items():
        np.save(args.output_dir / f"{model}_{target}_区间预测.npy", values)

    save_baseline_pipeline(
        args.output_dir / "baseline_interval_predictions.xlsx",
        dates,
        labels,
        predictions[("Baseline", "负载")],
        predictions[("Baseline", "光伏")],
    )
    all_path = args.output_dir / "全部模型区间预测与指标.xlsx"
    save_all_predictions(all_path, dates, labels, predictions, metrics)
    validate_excel(all_path, 1 + len(predictions))
    print(metrics.to_string(index=False))
    print(f"已生成：{all_path}")


if __name__ == "__main__":
    main()

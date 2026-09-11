"""问题二基准预测：负载取上周同期，光伏取此前 7 天同期均值。"""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


LOAD_SHEET = "小区负载"
PV_SHEET = "光伏发电实际功率"
FORECAST_START = datetime(2025, 2, 1)


def read_sheet(path: Path, sheet_name: str) -> tuple[list[datetime], list[object], np.ndarray]:
    """读取一个 365×144 的功率矩阵，并严格核对表结构。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()

    if len(rows) != 366 or len(rows[0]) != 145:
        raise ValueError(
            f"{sheet_name} 结构异常：期望 366 行×145 列，实际 {len(rows)} 行×{len(rows[0])} 列"
        )

    times = list(rows[0][1:])
    dates = [row[0] for row in rows[1:]]
    values = np.asarray([row[1:] for row in rows[1:]], dtype=float)

    if len(set(dates)) != len(dates):
        raise ValueError(f"{sheet_name} 存在重复日期")
    if np.isnan(values).any():
        raise ValueError(f"{sheet_name} 存在缺失或非数值功率")
    if (values < 0).any():
        raise ValueError(f"{sheet_name} 存在负功率")
    return dates, times, values


def make_baseline(
    dates: list[datetime], load: np.ndarray, pv: np.ndarray
) -> tuple[list[datetime], np.ndarray, np.ndarray]:
    """仅用预测日 0:00 前的数据生成逐日预测。"""
    first = next(i for i, date in enumerate(dates) if date >= FORECAST_START)
    forecast_dates = dates[first:]

    # 负载：严格采用前 7 天同一时段。
    load_prediction = load[first - 7 : len(dates) - 7].copy()

    # 光伏：对预测日前 7 天的同一时段取算术平均。
    pv_prediction = np.vstack([pv[i - 7 : i].mean(axis=0) for i in range(first, len(dates))])
    pv_prediction = np.clip(pv_prediction, 0.0, None)

    expected_shape = (len(forecast_dates), 144)
    if load_prediction.shape != expected_shape or pv_prediction.shape != expected_shape:
        raise AssertionError("基准预测矩阵维度不正确")
    return forecast_dates, load_prediction, pv_prediction


def style_header(worksheet) -> None:
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "B2"
    worksheet.auto_filter.ref = worksheet.dimensions


def write_wide_sheet(workbook: Workbook, name: str, dates, times, values) -> None:
    worksheet = workbook.create_sheet(name)
    worksheet.append(["日期\\时间", *times])
    for date, row in zip(dates, values, strict=True):
        worksheet.append([date, *[float(value) for value in row]])
    style_header(worksheet)
    worksheet.column_dimensions["A"].width = 13
    for column in range(2, 146):
        worksheet.column_dimensions[get_column_letter(column)].width = 10
    for cell in worksheet["A"][1:]:
        cell.number_format = "yyyy-mm-dd"
    for row in worksheet.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = "0.0000"


def write_output(path: Path, dates, times, load_prediction, pv_prediction) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    write_wide_sheet(workbook, "小区负载预测", dates, times, load_prediction)
    write_wide_sheet(workbook, "光伏发电预测", dates, times, pv_prediction)

    worksheet = workbook.create_sheet("预测长表")
    worksheet.append(["日期", "时段序号", "时刻", "负载预测功率", "光伏预测功率"])
    for day_index, date in enumerate(dates):
        for period_index, time_label in enumerate(times):
            worksheet.append(
                [
                    date,
                    period_index + 1,
                    time_label,
                    float(load_prediction[day_index, period_index]),
                    float(pv_prediction[day_index, period_index]),
                ]
            )
    style_header(worksheet)
    worksheet.column_dimensions["A"].width = 13
    worksheet.column_dimensions["B"].width = 10
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 18
    worksheet.column_dimensions["E"].width = 18
    for cell in worksheet["A"][1:]:
        cell.number_format = "yyyy-mm-dd"
    for cell in worksheet["C"][1:]:
        cell.number_format = "hh:mm"
    for column in ("D", "E"):
        for cell in worksheet[column][1:]:
            cell.number_format = "0.0000"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def validate_output(path: Path, expected_dates: int) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    expected = {
        "小区负载预测": (expected_dates + 1, 145),
        "光伏发电预测": (expected_dates + 1, 145),
        "预测长表": (expected_dates * 144 + 1, 5),
    }
    for sheet_name, shape in expected.items():
        worksheet = workbook[sheet_name]
        actual = (worksheet.max_row, worksheet.max_column)
        if actual != shape:
            raise AssertionError(f"{sheet_name} 维度错误：期望 {shape}，实际 {actual}")
    first_date = workbook["预测长表"]["A2"].value
    last_date = workbook["预测长表"][f"A{expected_dates * 144 + 1}"].value
    if first_date != datetime(2025, 2, 1) or last_date != datetime(2025, 12, 31):
        raise AssertionError(f"日期范围错误：{first_date} 至 {last_date}")
    workbook.close()


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input",
        type=Path,
        default=script_dir.parent / "附件2.xlsx",
        help="附件2.xlsx 路径",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=script_dir / "baseline_predictions.xlsx",
        help="输出 Excel 路径",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    load_dates, times, load = read_sheet(args.input, LOAD_SHEET)
    pv_dates, pv_times, pv = read_sheet(args.input, PV_SHEET)
    if load_dates != pv_dates or times != pv_times:
        raise ValueError("负载与光伏工作表的日期或时段不一致")

    dates, load_prediction, pv_prediction = make_baseline(load_dates, load, pv)
    write_output(args.output, dates, times, load_prediction, pv_prediction)
    validate_output(args.output, len(dates))

    print(f"已生成：{args.output}")
    print(f"预测日期：{dates[0]:%Y-%m-%d} 至 {dates[-1]:%Y-%m-%d}，共 {len(dates)} 天")
    print(f"每日电力时段：{len(times)}；长表记录：{len(dates) * len(times)}")
    print(f"负载预测范围：{load_prediction.min():.4f} 至 {load_prediction.max():.4f}")
    print(f"光伏预测范围：{pv_prediction.min():.4f} 至 {pv_prediction.max():.4f}")


if __name__ == "__main__":
    main()

"""输出问题二标准化区间预测：XGBoost负载与自适应加权光伏。"""

from __future__ import annotations

import argparse
from copy import copy
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from forecast_models import PV_SHEET, read_power_sheet


START_DATE = datetime(2025, 2, 1)
EXPECTED_DAYS = 334
EXPECTED_ENDPOINTS = 144


def format_endpoint(total_minutes: int) -> str:
    if total_minutes < 1440:
        hour, minute = divmod(total_minutes, 60)
        return f"{hour:02d}:{minute:02d}"
    hour, minute = divmod(total_minutes - 1440, 60)
    return f"次日{hour:02d}:{minute:02d}"


def interval_headers() -> list[str]:
    """00:10—00:20 至 次日00:00—次日00:10，共144个区间。"""
    return [
        f"{format_endpoint(10 + 10 * i)}—{format_endpoint(20 + 10 * i)}"
        for i in range(144)
    ]


def load_dates(reference_path: Path) -> list[datetime]:
    workbook = load_workbook(reference_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    dates = [worksheet.cell(row, 1).value for row in range(2, worksheet.max_row + 1)]
    workbook.close()
    if len(dates) != EXPECTED_DAYS or dates[0] != START_DATE or dates[-1] != datetime(2025, 12, 31):
        raise ValueError("参考文件日期范围不是 2025-02-01 至 2025-12-31 的334天")
    return dates


def load_predictions(
    question2_dir: Path, attachment2: Path, dates: list[datetime]
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    load_endpoints = np.load(question2_dir / "results" / "XGBoost_负载_predictions.npy")
    pv_endpoints = np.load(
        question2_dir / "improved_pv_results" / "自适应加权Baseline_predictions.npy"
    )
    if load_endpoints.shape != (EXPECTED_DAYS, EXPECTED_ENDPOINTS):
        raise ValueError(f"XGBoost负载预测维度异常：{load_endpoints.shape}")
    if pv_endpoints.shape != (EXPECTED_DAYS, EXPECTED_ENDPOINTS):
        raise ValueError(f"自适应光伏预测维度异常：{pv_endpoints.shape}")

    historical_dates, _, historical_pv = read_power_sheet(attachment2, PV_SHEET)
    load_dates, _, historical_load = read_power_sheet(attachment2, "小区负载")
    if load_dates != historical_dates:
        raise ValueError("附件2负载与光伏日期不一致")
    date_to_index = {date: i for i, date in enumerate(historical_dates)}
    planning_indices = np.asarray([date_to_index[date] for date in dates])

    # 目标为规划日次日00:10，其上周同期为规划日往前6天的00:10。
    load_extra_0010 = historical_load[planning_indices - 6, 0].astype(float, copy=True)

    # 规划时当天00:10尚不可得，采用规划日前7个已观测00:10的均值。
    pv_extra_0010 = np.asarray(
        [historical_pv[i - 7 : i, 0].mean() for i in planning_indices], dtype=float
    )
    return load_endpoints, load_extra_0010, pv_endpoints, pv_extra_0010


def to_interval_average(endpoints: np.ndarray, extra_0010: np.ndarray) -> np.ndarray:
    """144个日内端点加1个跨日端点，转换为144个区间平均值。"""
    within_day = (endpoints[:, :-1] + endpoints[:, 1:]) / 2
    final_interval = ((endpoints[:, -1] + extra_0010) / 2)[:, None]
    result = np.hstack([within_day, final_interval])
    if result.shape != (EXPECTED_DAYS, 144):
        raise AssertionError(f"区间矩阵维度异常：{result.shape}")
    if not np.isfinite(result).all() or np.any(result < 0):
        raise AssertionError("区间预测含缺失、无穷或负值")
    return result


def copy_reference_style(reference_path: Path, worksheet, rows: int, columns: int) -> None:
    reference_book = load_workbook(reference_path, read_only=False, data_only=False)
    reference_sheet = reference_book[reference_book.sheetnames[0]]
    # 参考文件的数值区使用默认 General 样式；只复制非默认的表头和日期列即可。
    def apply_style(source, target) -> None:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        target.number_format = source.number_format

    for column in range(1, columns + 1):
        source_column = min(column + 1, reference_sheet.max_column) if column > 1 else 1
        apply_style(reference_sheet.cell(1, source_column), worksheet.cell(1, column))
        source_letter = get_column_letter(source_column)
        target_letter = get_column_letter(column)
        worksheet.column_dimensions[target_letter].width = reference_sheet.column_dimensions[source_letter].width
    for row in range(2, rows + 1):
        apply_style(reference_sheet.cell(2, 1), worksheet.cell(row, 1))
    worksheet.row_dimensions[1].height = reference_sheet.row_dimensions[1].height
    reference_book.close()


def write_sheet(workbook: Workbook, name: str, dates, headers, values, reference_path: Path) -> None:
    worksheet = workbook.create_sheet(name)
    worksheet.append(["日期\\区间", *headers])
    for date, row in zip(dates, values, strict=True):
        worksheet.append([date, *[float(value) for value in row]])
    copy_reference_style(reference_path, worksheet, len(dates) + 1, len(headers) + 1)


def write_output(path: Path, reference_path: Path, dates, load_intervals, pv_intervals) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = interval_headers()
    write_sheet(workbook, "负载", dates, headers, load_intervals, reference_path)
    write_sheet(workbook, "光伏", dates, headers, pv_intervals, reference_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def validate_output(path: Path, load_expected: np.ndarray, pv_expected: np.ndarray) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if workbook.sheetnames != ["负载", "光伏"]:
        raise AssertionError(f"工作表异常：{workbook.sheetnames}")
    headers = interval_headers()
    for sheet_name, expected in (("负载", load_expected), ("光伏", pv_expected)):
        worksheet = workbook[sheet_name]
        if (worksheet.max_row, worksheet.max_column) != (335, 145):
            raise AssertionError(f"{sheet_name} 尺寸异常")
        row_iterator = worksheet.iter_rows(values_only=True)
        header_row = next(row_iterator)
        actual_headers = list(header_row[1:])
        if actual_headers != headers:
            raise AssertionError(f"{sheet_name} 区间表头异常")
        data_rows = list(row_iterator)
        actual_dates = [row[0] for row in data_rows]
        if actual_dates[0] != START_DATE or actual_dates[-1] != datetime(2025, 12, 31):
            raise AssertionError(f"{sheet_name} 日期范围异常")
        actual = np.asarray([row[1:] for row in data_rows], dtype=float)
        if not np.allclose(actual, expected, rtol=0, atol=1e-10):
            raise AssertionError(f"{sheet_name} 写入数值与计算结果不一致")
    workbook.close()


def main() -> None:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--reference",
        type=Path,
        default=script_dir.parent.parent / "演示数据用于参考格式.xlsx",
    )
    parser.add_argument("--attachment2", type=Path, default=script_dir.parent / "附件2.xlsx")
    parser.add_argument(
        "--output",
        type=Path,
        default=script_dir / "标准化负载与光伏区间预测.xlsx",
    )
    args = parser.parse_args()

    dates = load_dates(args.reference)
    load_endpoints, load_extra, pv_endpoints, pv_extra = load_predictions(
        script_dir, args.attachment2, dates
    )
    load_intervals = to_interval_average(load_endpoints, load_extra)
    pv_intervals = to_interval_average(pv_endpoints, pv_extra)
    write_output(args.output, args.reference, dates, load_intervals, pv_intervals)
    validate_output(args.output, load_intervals, pv_intervals)
    print(f"输出文件：{args.output}")
    print(f"工作表：负载、光伏；日期：{dates[0]:%Y-%m-%d} 至 {dates[-1]:%Y-%m-%d}")
    print(f"每个工作表：{len(dates)}天 × {load_intervals.shape[1]}个区间")
    print(f"负载区间范围：{load_intervals.min():.6f} 至 {load_intervals.max():.6f}")
    print(f"光伏区间范围：{pv_intervals.min():.6f} 至 {pv_intervals.max():.6f}")


if __name__ == "__main__":
    main()

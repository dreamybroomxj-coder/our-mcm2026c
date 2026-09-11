"""问题二：按日滚动预测负载与光伏（Baseline/XGBoost/LightGBM/SARIMA）。

所有预测都模拟当天 0:00 的信息边界：预测某日时只允许使用此前日期的真实值。
"""

from __future__ import annotations

import argparse
import json
import math
import time
import warnings
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable

import numpy as np
import pandas as pd
from lightgbm import LGBMRegressor
from openpyxl import load_workbook
from sklearn.metrics import mean_absolute_error, mean_squared_error
from statsmodels.tsa.statespace.sarimax import SARIMAX
from statsmodels.tools.sm_exceptions import ConvergenceWarning
from xgboost import XGBRegressor


SEED = 2026
LOAD_SHEET = "小区负载"
PV_SHEET = "光伏发电实际功率"
TARGETS = ("负载", "光伏")
DEFAULT_START = datetime(2025, 2, 1)
DEFAULT_END = datetime(2025, 12, 31)


@dataclass(frozen=True)
class Config:
    start_date: str
    end_date: str
    tree_refit_every: int = 1
    sarima_refit_every: int = 30
    n_estimators: int = 120
    max_depth: int = 5
    learning_rate: float = 0.05
    sarima_maxiter: int = 30
    seed: int = SEED


def read_power_sheet(path: Path, sheet_name: str) -> tuple[list[datetime], list[object], np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    if len(rows) != 366 or len(rows[0]) != 145:
        raise ValueError(f"{sheet_name} 应为 366×145，实际为 {len(rows)}×{len(rows[0])}")
    dates = [pd.Timestamp(row[0]).to_pydatetime() for row in rows[1:]]
    times = list(rows[0][1:])
    values = np.asarray([row[1:] for row in rows[1:]], dtype=float)
    if values.shape != (365, 144) or not np.isfinite(values).all():
        raise ValueError(f"{sheet_name} 数值矩阵异常")
    if np.any(values < 0):
        raise ValueError(f"{sheet_name} 含负功率")
    return dates, times, values


def load_data(path: Path):
    dates, times, load = read_power_sheet(path, LOAD_SHEET)
    pv_dates, pv_times, pv = read_power_sheet(path, PV_SHEET)
    if dates != pv_dates or times != pv_times:
        raise ValueError("两个工作表的日期或时段不一致")
    if dates[0] != datetime(2025, 1, 1) or dates[-1] != datetime(2025, 12, 31):
        raise ValueError(f"日期范围异常：{dates[0]} 至 {dates[-1]}")
    return dates, times, {"负载": load, "光伏": pv}


def date_indices(dates, start: datetime, end: datetime) -> np.ndarray:
    indices = np.asarray([i for i, date in enumerate(dates) if start <= date <= end], dtype=int)
    if len(indices) == 0:
        raise ValueError("指定预测日期范围没有数据")
    if indices[0] < 31:
        raise ValueError("预测起点必须不早于 2025-02-01，以保证至少一个月历史数据")
    return indices


def baseline_predict(values: np.ndarray, indices: np.ndarray, target: str) -> np.ndarray:
    if target == "负载":
        return np.vstack([values[i - 7] for i in indices])
    return np.vstack([values[i - 7 : i].mean(axis=0) for i in indices])


def calendar_features(date: datetime, slots: int = 144) -> np.ndarray:
    slot = np.arange(slots)
    dow = date.weekday()
    doy = date.timetuple().tm_yday
    return np.column_stack(
        [
            slot / (slots - 1),
            np.sin(2 * np.pi * slot / slots),
            np.cos(2 * np.pi * slot / slots),
            np.full(slots, dow / 6),
            np.full(slots, float(dow >= 5)),
            np.full(slots, np.sin(2 * np.pi * dow / 7)),
            np.full(slots, np.cos(2 * np.pi * dow / 7)),
            np.full(slots, np.sin(2 * np.pi * doy / 365)),
            np.full(slots, np.cos(2 * np.pi * doy / 365)),
        ]
    )


def day_features(values: np.ndarray, dates: list[datetime], day_index: int) -> np.ndarray:
    """构造在 day_index 当天 0:00 已知的特征，不读取当天真实值。"""
    if day_index < 14:
        raise ValueError("树模型特征至少需要 14 天历史")
    history_7 = values[day_index - 7 : day_index]
    history_14 = values[day_index - 14 : day_index]
    return np.column_stack(
        [
            calendar_features(dates[day_index], values.shape[1]),
            values[day_index - 1],
            values[day_index - 2],
            values[day_index - 7],
            values[day_index - 14],
            history_7.mean(axis=0),
            np.median(history_7, axis=0),
            history_7.std(axis=0),
            history_14.mean(axis=0),
            history_14.std(axis=0),
        ]
    )


def training_matrix(values: np.ndarray, dates: list[datetime], forecast_day: int):
    days = range(14, forecast_day)
    x = np.vstack([day_features(values, dates, i) for i in days])
    y = np.concatenate([values[i] for i in days])
    return x, y


def make_xgboost(config: Config):
    return XGBRegressor(
        objective="reg:squarederror",
        n_estimators=config.n_estimators,
        max_depth=config.max_depth,
        learning_rate=config.learning_rate,
        min_child_weight=8,
        subsample=0.85,
        colsample_bytree=0.85,
        reg_lambda=2.0,
        random_state=config.seed,
        n_jobs=-1,
        tree_method="hist",
    )


def make_lightgbm(config: Config):
    return LGBMRegressor(
        objective="regression",
        n_estimators=config.n_estimators,
        max_depth=config.max_depth,
        num_leaves=min(31, 2**config.max_depth),
        learning_rate=config.learning_rate,
        min_child_samples=30,
        subsample=0.85,
        colsample_bytree=0.85,
        reg_lambda=2.0,
        random_state=config.seed,
        n_jobs=-1,
        verbosity=-1,
    )


def tree_predict(
    values: np.ndarray,
    dates: list[datetime],
    indices: np.ndarray,
    model_factory: Callable[[], object],
    refit_every: int,
) -> np.ndarray:
    predictions = []
    model = None
    for offset, day_index in enumerate(indices):
        if model is None or offset % refit_every == 0:
            x_train, y_train = training_matrix(values, dates, day_index)
            model = model_factory()
            model.fit(x_train, y_train)
        pred = np.asarray(model.predict(day_features(values, dates, day_index)), dtype=float)
        predictions.append(np.clip(pred, 0.0, None))
    return np.vstack(predictions)


def fit_sarima(endog: np.ndarray, maxiter: int):
    """对每日总功率和建立带 7 日季节周期的 SARIMA。"""
    model = SARIMAX(
        endog,
        order=(1, 0, 1),
        seasonal_order=(1, 0, 0, 7),
        trend="c",
        enforce_stationarity=False,
        enforce_invertibility=False,
        simple_differencing=False,
    )
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", ConvergenceWarning)
        result = model.fit(disp=False, maxiter=maxiter, method="lbfgs")
    if not result.mle_retvals.get("converged", False):
        # 早期样本较少时 L-BFGS 偶尔停在迭代上限，改用无梯度方法复核。
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", ConvergenceWarning)
            result = model.fit(disp=False, maxiter=max(100, maxiter), method="powell")
    if not result.mle_retvals.get("converged", False):
        raise RuntimeError("SARIMA 参数估计未收敛")
    return result


def normalized_profile(values: np.ndarray, day_index: int, target: str) -> np.ndarray:
    """只用历史日曲线，将 SARIMA 的日总量分配到 144 个时段。"""
    if target == "负载":
        # 优先使用过去四个同星期日；数据不足时退回最近七天。
        candidates = [i for i in range(day_index - 7, -1, -7)][:4]
        history = values[candidates] if candidates else values[day_index - 7 : day_index]
    else:
        history = values[day_index - 7 : day_index]
    totals = history.sum(axis=1)
    valid = totals > 1e-8
    if not valid.any():
        return np.full(values.shape[1], 1.0 / values.shape[1])
    profiles = history[valid] / totals[valid, None]
    profile = profiles.mean(axis=0)
    return profile / profile.sum()


def sarima_predict(
    values: np.ndarray,
    indices: np.ndarray,
    refit_every: int,
    maxiter: int,
    target: str,
) -> np.ndarray:
    daily_totals = values.sum(axis=1)
    predictions = []
    result = None
    state_end_day = None
    for offset, day_index in enumerate(indices):
        if result is None or offset % refit_every == 0:
            result = fit_sarima(daily_totals[:day_index], maxiter)
            state_end_day = day_index
        elif state_end_day != day_index:
            raise AssertionError("SARIMA 状态更新时间不连续")

        predicted_total = max(0.0, float(np.asarray(result.forecast(steps=1))[0]))
        profile = normalized_profile(values, day_index, target)
        predictions.append(predicted_total * profile)

        # 当天结束后才将当天真实值加入状态，供下一日 0:00 使用。
        result = result.append([daily_totals[day_index]], refit=False)
        state_end_day = day_index + 1
    return np.vstack(predictions)


def metrics(actual: np.ndarray, predicted: np.ndarray, target: str) -> dict[str, float]:
    error = predicted - actual
    denominator = np.abs(actual).sum()
    result = {
        "MAE": float(mean_absolute_error(actual.ravel(), predicted.ravel())),
        "RMSE": float(math.sqrt(mean_squared_error(actual.ravel(), predicted.ravel()))),
        "WAPE_percent": float(np.abs(error).sum() / denominator * 100),
        "Bias": float(error.mean()),
    }
    if target == "光伏":
        daylight = actual > 1.0
        result["Daylight_MAE"] = float(np.abs(error[daylight]).mean())
        result["Daylight_WAPE_percent"] = float(
            np.abs(error[daylight]).sum() / np.abs(actual[daylight]).sum() * 100
        )
    return result


def evaluate(predictions, actuals, dates, indices) -> tuple[pd.DataFrame, pd.DataFrame]:
    overall_rows = []
    monthly_rows = []
    for model_name, target_predictions in predictions.items():
        for target in TARGETS:
            actual = actuals[target][indices]
            pred = target_predictions[target]
            overall_rows.append({"模型": model_name, "对象": target, **metrics(actual, pred, target)})
            months = sorted({dates[i].month for i in indices})
            for month in months:
                mask = np.asarray([dates[i].month == month for i in indices])
                monthly_rows.append(
                    {"模型": model_name, "对象": target, "月份": month, **metrics(actual[mask], pred[mask], target)}
                )
    return pd.DataFrame(overall_rows), pd.DataFrame(monthly_rows)


def save_predictions(path: Path, dates, times, indices, predictions, actuals) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        rows = []
        for model_name, target_predictions in predictions.items():
            for local_day, day_index in enumerate(indices):
                for slot in range(144):
                    rows.append(
                        {
                            "模型": model_name,
                            "日期": dates[day_index],
                            "时段序号": slot + 1,
                            "时刻": times[slot],
                            "负载预测功率": target_predictions["负载"][local_day, slot],
                            "光伏预测功率": target_predictions["光伏"][local_day, slot],
                            "负载实际功率": actuals["负载"][day_index, slot],
                            "光伏实际功率": actuals["光伏"][day_index, slot],
                        }
                    )
        pd.DataFrame(rows).to_excel(writer, sheet_name="全部预测长表", index=False)
        for model_name, target_predictions in predictions.items():
            for target in TARGETS:
                frame = pd.DataFrame(target_predictions[target], columns=[str(t) for t in times])
                frame.insert(0, "日期\\时间", [dates[i] for i in indices])
                sheet = f"{model_name}_{target}"[:31]
                frame.to_excel(writer, sheet_name=sheet, index=False)


def parse_methods(text: str) -> list[str]:
    allowed = {"Baseline", "XGBoost", "LightGBM", "SARIMA"}
    methods = [item.strip() for item in text.split(",") if item.strip()]
    unknown = set(methods) - allowed
    if unknown:
        raise ValueError(f"未知模型：{sorted(unknown)}")
    return methods


def main() -> None:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=script_dir.parent / "附件2.xlsx")
    parser.add_argument("--output-dir", type=Path, default=script_dir / "results")
    parser.add_argument("--start-date", default="2025-02-01")
    parser.add_argument("--end-date", default="2025-12-31")
    parser.add_argument("--methods", default="Baseline,XGBoost,LightGBM,SARIMA")
    parser.add_argument("--tree-refit-every", type=int, default=1)
    parser.add_argument("--sarima-refit-every", type=int, default=30)
    parser.add_argument("--n-estimators", type=int, default=120)
    parser.add_argument("--sarima-maxiter", type=int, default=30)
    args = parser.parse_args()

    config = Config(
        start_date=args.start_date,
        end_date=args.end_date,
        tree_refit_every=args.tree_refit_every,
        sarima_refit_every=args.sarima_refit_every,
        n_estimators=args.n_estimators,
        max_depth=5,
        learning_rate=0.05,
        sarima_maxiter=args.sarima_maxiter,
    )
    if min(config.tree_refit_every, config.sarima_refit_every, config.n_estimators) <= 0:
        raise ValueError("重训练间隔和树数量必须为正整数")

    dates, times, actuals = load_data(args.input)
    indices = date_indices(dates, datetime.fromisoformat(args.start_date), datetime.fromisoformat(args.end_date))
    methods = parse_methods(args.methods)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    predictions: dict[str, dict[str, np.ndarray]] = {}

    for method in methods:
        started = time.perf_counter()
        predictions[method] = {}
        for target in TARGETS:
            values = actuals[target]
            if method == "Baseline":
                pred = baseline_predict(values, indices, target)
            elif method == "XGBoost":
                pred = tree_predict(values, dates, indices, lambda: make_xgboost(config), config.tree_refit_every)
            elif method == "LightGBM":
                pred = tree_predict(values, dates, indices, lambda: make_lightgbm(config), config.tree_refit_every)
            else:
                pred = sarima_predict(
                    values, indices, config.sarima_refit_every, config.sarima_maxiter, target
                )
            predictions[method][target] = pred
            np.save(args.output_dir / f"{method}_{target}_predictions.npy", pred)
        print(f"{method} 完成，耗时 {time.perf_counter() - started:.1f} 秒")

    overall, monthly = evaluate(predictions, actuals, dates, indices)
    overall.to_csv(args.output_dir / "overall_metrics.csv", index=False, encoding="utf-8-sig")
    monthly.to_csv(args.output_dir / "monthly_metrics.csv", index=False, encoding="utf-8-sig")
    save_predictions(args.output_dir / "all_model_predictions.xlsx", dates, times, indices, predictions, actuals)
    manifest = {
        "input": str(args.input.resolve()),
        "config": asdict(config),
        "methods": methods,
        "prediction_days": len(indices),
        "periods_per_day": 144,
        "information_boundary": "预测某日时仅使用此前日期的真实数据",
    }
    (args.output_dir / "run_config.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(overall.to_string(index=False))


if __name__ == "__main__":
    main()

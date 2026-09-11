"""输出问题二最终采用模型的 144 个十分钟区间平均预测。"""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from forecast_models import LOAD_SHEET, PV_SHEET, read_power_sheet


START_INDEX = 31  # 2025-02-01
EXPECTED_DAYS = 334


def interval_headers() -> list[str]:
    """生成 00:10—00:20 至次日 00:00—00:10 的 144 个区间。"""
    headers = []
    for interval_index in range(144):
        start_minutes = 10 + interval_index * 10
        end_minutes = start_minutes + 10

        def label(total_minutes: int) -> str:
            if total_minutes < 1440:
                return f"{total_minutes // 60:02d}:{total_minutes % 60:02d}"
            shifted = total_minutes - 1440
            return f"次日{shifted // 60:02d}:{shifted % 60:02d}"

        headers.append(f"{label(start_minutes)}—{label(end_minutes)}")
    if headers[0] != "00:10—00:20" or headers[-1] != "次日00:00—次日00:10":
        raise AssertionError("区间表头生成错误")
    return headers


def next_0010_load_weekly(load_actual: np.ndarray, day_indices: np.ndarray) -> np.ndarray:
    """目标 d+1 日 00:10 的上周同期为 d-6 日 00:10。"""
    return np.asarray([load_actual[day_index - 6, 0] for day_index in day_indices], dtype=float)


def next_0010_pv_history_mean(pv_actual: np.ndarray, day_indices: np.ndarray) -> np.ndarray:
    """使用规划时刻已观测的 d-7 至 d-1 共 7 个 00:10 功率。"""
    return np.asarray(
        [pv_actual[day_index - 7 : day_index, 0].mean() for day_index in day_indices],
        dtype=float,
    )


def endpoints_to_intervals(main_prediction: np.ndarray, final_endpoint: np.ndarray) -> np.ndarray:
    if main_prediction.shape != (EXPECTED_DAYS, 144) or final_endpoint.shape != (EXPECTED_DAYS,):
        raise ValueError(
            f"预测维度异常：主预测 {main_prediction.shape}，跨日端点 {final_endpoint.shape}"
        )
    endpoints = np.column_stack([main_prediction, final_endpoint])
    intervals = (endpoints[:, :-1] + endpoints[:, 1:]) / 2.0
    if intervals.shape != (EXPECTED_DAYS, 144):
        raise AssertionError("区间平均矩阵维度错误")
    if not np.isfinite(intervals).all() or np.any(intervals < 0):
        raise ValueError("区间平均预测含缺失、无穷或负值")
    return intervals


def write_sheet(workbook: Workbook, name: str, dates, headers, values) -> None:
    worksheet = workbook.create_sheet(name)
    worksheet.append(["日期\\区间", *headers])
    for date, row in zip(dates, values, strict=True):
        worksheet.append([date, *[float(value) for value in row]])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "B2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 13
    for column_index in range(2, 146):
        worksheet.column_dimensions[get_column_letter(column_index)].width = 19
    for cell in worksheet["A"][1:]:
        cell.number_format = "yyyy-mm-dd"
        cell.alignment = Alignment(horizontal="center")
    for row in worksheet.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = "0.0000"


def validate_output(path: Path, dates: list[datetime], headers: list[str]) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    if workbook.sheetnames != ["XGBoost_负载", "自适应加权_光伏"]:
        raise AssertionError(f"工作表名称错误：{workbook.sheetnames}")
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        if (worksheet.max_row, worksheet.max_column) != (335, 145):
            raise AssertionError(
                f"{sheet_name} 结构错误：{worksheet.max_row} 行×{worksheet.max_column} 列"
            )
        actual_headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        if actual_headers != ["日期\\区间", *headers]:
            raise AssertionError(f"{sheet_name} 表头不一致")
        if worksheet.cell(2, 1).value != dates[0] or worksheet.cell(335, 1).value != dates[-1]:
            raise AssertionError(f"{sheet_name} 日期范围错误")
        values = np.asarray(
            list(worksheet.iter_rows(min_row=2, min_col=2, values_only=True)), dtype=float
        )
        if values.shape != (334, 144) or not np.isfinite(values).all() or np.any(values < 0):
            raise AssertionError(f"{sheet_name} 数值检查失败")
    workbook.close()


def main() -> None:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=script_dir.parent / "附件2.xlsx")
    parser.add_argument("--load-prediction", type=Path, default=script_dir / "results" / "XGBoost_负载_predictions.npy")
    parser.add_argument(
        "--pv-prediction",
        type=Path,
        default=script_dir / "improved_pv_results" / "自适应加权Baseline_predictions.npy",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=script_dir / "问题二_标准化区间预测.xlsx",
    )
    args = parser.parse_args()

    load_dates, _, load_actual = read_power_sheet(args.input, LOAD_SHEET)
    pv_dates, _, pv_actual = read_power_sheet(args.input, PV_SHEET)
    if load_dates != pv_dates:
        raise ValueError("负载和光伏日期不一致")
    dates = load_dates[START_INDEX:]
    day_indices = np.arange(START_INDEX, len(load_dates))
    if len(dates) != EXPECTED_DAYS or dates[0] != datetime(2025, 2, 1) or dates[-1] != datetime(2025, 12, 31):
        raise ValueError("输出日期范围异常")

    load_main = np.load(args.load_prediction)
    pv_main = np.load(args.pv_prediction)
    load_intervals = endpoints_to_intervals(
        load_main, next_0010_load_weekly(load_actual, day_indices)
    )
    pv_intervals = endpoints_to_intervals(
        pv_main, next_0010_pv_history_mean(pv_actual, day_indices)
    )

    headers = interval_headers()
    workbook = Workbook()
    workbook.remove(workbook.active)
    write_sheet(workbook, "XGBoost_负载", dates, headers, load_intervals)
    write_sheet(workbook, "自适应加权_光伏", dates, headers, pv_intervals)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    validate_output(args.output, dates, headers)

    print(f"已生成：{args.output}")
    print(f"日期：{dates[0]:%Y-%m-%d} 至 {dates[-1]:%Y-%m-%d}，共 {len(dates)} 天")
    print(f"每个工作表：{len(dates)} 行数据×{len(headers)} 个区间")
    print(f"负载区间范围：{load_intervals.min():.4f} 至 {load_intervals.max():.4f}")
    print(f"光伏区间范围：{pv_intervals.min():.4f} 至 {pv_intervals.max():.4f}")


if __name__ == "__main__":
    main()

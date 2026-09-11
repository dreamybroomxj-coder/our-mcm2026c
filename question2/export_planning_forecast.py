"""导出最终选定预测模型的数据，供问题二规划模型直接读取。

最终模型：负载使用 XGBoost，光伏使用自适应加权 Baseline。
所有功率均为 10 分钟区间平均功率，而非瞬时端点值。
"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from interval_preprocessing import attachment1_tables, period_mapping


FORECAST_START = datetime(2025, 2, 1)
FORECAST_END = datetime(2025, 12, 31)
EXPECTED_DAYS = 334
PERIODS_PER_DAY = 144
LOAD_MODEL = "XGBoost"
PV_MODEL = "自适应加权Baseline"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def forecast_dates() -> list[datetime]:
    dates = []
    current = FORECAST_START
    while current <= FORECAST_END:
        dates.append(current)
        current += timedelta(days=1)
    if len(dates) != EXPECTED_DAYS:
        raise AssertionError(f"预测日期数量异常：{len(dates)}")
    return dates


def read_prediction(path: Path, label: str) -> np.ndarray:
    values = np.load(path)
    expected = (EXPECTED_DAYS, PERIODS_PER_DAY)
    if values.shape != expected:
        raise ValueError(f"{label}维度异常：期望{expected}，实际{values.shape}")
    if not np.isfinite(values).all():
        raise ValueError(f"{label}存在NaN或无穷值")
    if np.any(values < 0):
        raise ValueError(f"{label}存在负功率")
    return values.astype(float, copy=False)


def read_interval_price(attachment1: Path) -> np.ndarray:
    _, interval_table = attachment1_tables(attachment1)
    # attachment1_tables 的第7列为区间平均电价。
    prices = interval_table.iloc[:, 6].to_numpy(dtype=float)
    if prices.shape != (PERIODS_PER_DAY,) or not np.isfinite(prices).all():
        raise ValueError("附件1区间平均电价异常")
    if np.any(prices < 0):
        raise ValueError("附件1区间平均电价存在负值")
    return prices


def build_long_table(
    dates: list[datetime],
    mapping: pd.DataFrame,
    prices: np.ndarray,
    load_prediction: np.ndarray,
    pv_prediction: np.ndarray,
) -> pd.DataFrame:
    net_load = load_prediction - pv_prediction
    rows = EXPECTED_DAYS * PERIODS_PER_DAY
    table = pd.DataFrame(
        {
            "日期": np.repeat(dates, PERIODS_PER_DAY),
            "时段序号": np.tile(np.arange(1, PERIODS_PER_DAY + 1), EXPECTED_DAYS),
            "开始时刻": np.tile(mapping["开始时刻"].to_numpy(), EXPECTED_DAYS),
            "结束时刻": np.tile(mapping["结束时刻"].to_numpy(), EXPECTED_DAYS),
            "时间区间": np.tile(mapping["时间区间"].to_numpy(), EXPECTED_DAYS),
            "时长_h": np.full(rows, 1 / 6),
            "区间平均电价": np.tile(prices, EXPECTED_DAYS),
            "预测负载功率": load_prediction.ravel(),
            "预测光伏功率": pv_prediction.ravel(),
            "预测净负荷功率": net_load.ravel(),
        }
    )
    return table


def wide_table(dates: list[datetime], labels: list[str], values: np.ndarray) -> pd.DataFrame:
    frame = pd.DataFrame(values, columns=labels)
    frame.insert(0, "日期\\区间", dates)
    return frame


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in worksheet.columns:
            sample = list(column)[:200]
            width = min(24, max(10, max(len(str(cell.value or "")) for cell in sample) + 2))
            worksheet.column_dimensions[column[0].column_letter].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    workbook.save(path)


def export(
    load_path: Path,
    pv_path: Path,
    attachment1: Path,
    output_dir: Path,
) -> None:
    dates = forecast_dates()
    mapping = period_mapping()
    labels = mapping["时间区间"].tolist()
    load_prediction_data = read_prediction(load_path, "XGBoost负载预测")
    pv_prediction = read_prediction(pv_path, "自适应加权光伏预测")
    prices = read_interval_price(attachment1)
    net_load = load_prediction_data - pv_prediction
    long_table = build_long_table(
        dates, mapping, prices, load_prediction_data, pv_prediction
    )

    output_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = output_dir / "规划模型预测输入.xlsx"
    csv_path = output_dir / "规划模型预测输入.csv"
    npz_path = output_dir / "规划模型预测输入.npz"

    notes = pd.DataFrame(
        [
            ["预测日期", "2025-02-01至2025-12-31，共334天"],
            ["时段数量", "每天144个10分钟区间"],
            ["负载模型", LOAD_MODEL],
            ["光伏模型", PV_MODEL],
            ["功率口径", "相邻瞬时端点平均得到的10分钟区间平均功率"],
            ["00:00边界", "规划时刻真实值已在端点转区间过程中使用"],
            ["净负荷定义", "预测负载功率-预测光伏功率；允许为负，表示预测光伏盈余"],
            ["电价口径", "附件1相邻端点平均；按题意每天采用同一日内价格序列"],
            ["规划主输入", "工作表“规划模型输入”或同名CSV"],
        ],
        columns=["项目", "说明"],
    )
    price_table = mapping[["时段序号", "开始时刻", "结束时刻", "时间区间", "时长/h"]].copy()
    price_table["区间平均电价"] = prices

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        notes.to_excel(writer, sheet_name="使用说明", index=False)
        long_table.to_excel(writer, sheet_name="规划模型输入", index=False)
        wide_table(dates, labels, load_prediction_data).to_excel(
            writer, sheet_name="负载预测_宽表", index=False
        )
        wide_table(dates, labels, pv_prediction).to_excel(
            writer, sheet_name="光伏预测_宽表", index=False
        )
        wide_table(dates, labels, net_load).to_excel(
            writer, sheet_name="净负荷预测_宽表", index=False
        )
        price_table.to_excel(writer, sheet_name="日内区间电价", index=False)
    style_workbook(xlsx_path)

    long_table.to_csv(csv_path, index=False, encoding="utf-8-sig", float_format="%.6f")
    np.savez_compressed(
        npz_path,
        dates=np.asarray([date.strftime("%Y-%m-%d") for date in dates]),
        period=np.arange(1, PERIODS_PER_DAY + 1),
        interval_labels=np.asarray(labels),
        price=prices,
        load_prediction=load_prediction_data,
        pv_prediction=pv_prediction,
        net_load_prediction=net_load,
        duration_hours=np.asarray(1 / 6),
    )

    manifest = {
        "load_model": LOAD_MODEL,
        "pv_model": PV_MODEL,
        "prediction_period": ["2025-02-01", "2025-12-31"],
        "days": EXPECTED_DAYS,
        "periods_per_day": PERIODS_PER_DAY,
        "records": len(long_table),
        "load_input": {"path": str(load_path.resolve()), "sha256": sha256(load_path)},
        "pv_input": {"path": str(pv_path.resolve()), "sha256": sha256(pv_path)},
        "price_input": {"path": str(attachment1.resolve()), "sha256": sha256(attachment1)},
        "outputs": [xlsx_path.name, csv_path.name, npz_path.name],
    }
    (output_dir / "导出清单.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    validate_outputs(xlsx_path, csv_path, npz_path, load_prediction_data, pv_prediction)


def validate_outputs(
    xlsx_path: Path,
    csv_path: Path,
    npz_path: Path,
    expected_load: np.ndarray,
    expected_pv: np.ndarray,
) -> None:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    expected_sheets = [
        "使用说明",
        "规划模型输入",
        "负载预测_宽表",
        "光伏预测_宽表",
        "净负荷预测_宽表",
        "日内区间电价",
    ]
    if workbook.sheetnames != expected_sheets:
        raise AssertionError(f"工作表结构异常：{workbook.sheetnames}")
    if workbook["规划模型输入"].max_row != EXPECTED_DAYS * PERIODS_PER_DAY + 1:
        raise AssertionError("规划模型输入行数异常")
    if workbook["规划模型输入"].max_column != 10:
        raise AssertionError("规划模型输入列数异常")
    workbook.close()

    csv_table = pd.read_csv(csv_path, header=0)
    if len(csv_table) != EXPECTED_DAYS * PERIODS_PER_DAY:
        raise AssertionError("CSV记录数异常")
    if not np.allclose(csv_table["预测负载功率"].to_numpy(), expected_load.ravel(), atol=5e-7):
        raise AssertionError("CSV负载预测与源数组不一致")
    if not np.allclose(csv_table["预测光伏功率"].to_numpy(), expected_pv.ravel(), atol=5e-7):
        raise AssertionError("CSV光伏预测与源数组不一致")

    data = np.load(npz_path)
    if not np.array_equal(data["load_prediction"], expected_load):
        raise AssertionError("NPZ负载预测与源数组不一致")
    if not np.array_equal(data["pv_prediction"], expected_pv):
        raise AssertionError("NPZ光伏预测与源数组不一致")


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--load",
        type=Path,
        default=script_dir / "interval_results" / "XGBoost_负载_区间预测.npy",
    )
    parser.add_argument(
        "--pv",
        type=Path,
        default=script_dir / "interval_results" / "自适应加权Baseline_光伏_区间预测.npy",
    )
    parser.add_argument("--attachment1", type=Path, default=script_dir.parent / "附件1.xlsx")
    parser.add_argument("--output-dir", type=Path, default=script_dir / "planning_input")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    export(args.load, args.pv, args.attachment1, args.output_dir)
    print(f"已生成：{args.output_dir / '规划模型预测输入.xlsx'}")
    print(f"规划记录数：{EXPECTED_DAYS * PERIODS_PER_DAY}")
    print(f"负载模型：{LOAD_MODEL}；光伏模型：{PV_MODEL}")


if __name__ == "__main__":
    main()
